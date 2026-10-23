"""Разовый перенос ответов старой Google-формы в УЖЕ существующие карточки «Карты команды» (team_people, team_photos).

Запуск (разовая команда, к HTTP не подключается; в контейнере API, под пользователем hub_app):

  python -m app.modules.team.import_responses DIR [--aliases aliases.json]            # сухой прогон (по умолчанию)
  python -m app.modules.team.import_responses DIR [--aliases aliases.json] --apply    # запись

DIR: responses.xlsx (лист «Ответы на форму»), emails.xlsx (лист Sheet1: «Имя» латиницей, «Почта»), каталог photos/ с файлами
<driveId>.full.jpg и <driveId>.thumb.jpg. aliases.json: {"Фамилия Имя из ответа": "Фамилия Имя карточки"}.
Пишется только пустое: имя/фамилия/должность, руководитель, направление, тег, команда, show_in_map, user_id не меняются.
Всё — одна транзакция, сухой прогон делает ROLLBACK; повторный запуск безопасен.
Коды выхода: 0 — успех / чистый сухой прогон; 1 — ошибки входных файлов; 2 — конфликт с БД.
Персональные данные (имена, почты, telegram, тексты) не печатаются и не пишутся в журнал: только счётчики и номера строк Excel.
"""
import argparse
import itertools
import json
import re
import sys
from difflib import SequenceMatcher
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import text
from sqlalchemy.exc import DBAPIError
from sqlalchemy.orm import Session

from app.core.audit import log_action
from app.db.database import get_engine
from app.modules.team import service
from app.modules.team.seed import LOCK_TEAM, Report

ACTION = "team.import_responses"
S = service.S

CHIP_MAX = service.CHIP_LIMIT
CHIPS_MAX = 20
EMAIL_MAX = 120
WORD_MIN = 0.8  # каждое слово при сопоставлении почты
LEAD_MIN = 0.1  # отрыв лидера от второго кандидата
MAX_WORDS = 4

# Префиксы заголовков -> ключ. Обязательные столбцы responses.xlsx.
RESPONSE_COLS = {
    "first": "имя", "last": "фамилия", "responsibility": "расскажи, какая у тебя сейчас зона", "ask_me_about": "приведи примеры",
    "location": "в какой локации", "work_format": "какой у тебя формат", "telegram": "как тебя найти в телеграм",
    "hobbies": "какие еще у тебя есть таланты", "photo": "закинь сюда свою фотку", "email": "напиши свою корпоративную почту",
}
EMAIL_COLS = {"name": "имя", "email": "почта"}
TG_PREFIX = re.compile(r"^(?:https?://)?(?:www\.)?(?:t\.me|telegram\.me)/", re.I)
TG_RE = re.compile(r"^[A-Za-z0-9_]{3,64}$")
DRIVE_ID = re.compile(r"[?&]id=([A-Za-z0-9_-]+)")
DRIVE_PATH_ID = re.compile(r"/d/([A-Za-z0-9_-]+)")

REASONS = {
    "empty_name": "ответ без имени/фамилии",
    "not_found": "нет подходящей карточки",
    "ambiguous": "несколько подходящих карточек (нужен алиас)",
    "alias_not_found": "алиас: карточка не найдена",
    "alias_ambiguous": "алиас: несколько карточек",
    "telegram": "telegram недопустим",
    "too_long": "значение слишком длинное (поле пропущено)",
    "chips_cut": "чипов больше 20 (лишние отброшены)",
    "email_bad": "почта в ответе некорректна",
    "email_dup": "почта уже занята другой карточкой (пропуск)",
    "photo_no_link": "фото: нет ссылки/ID",
    "photo_missing": "фото: нет файла full/thumb",
    "photo_bad": "фото: не JPEG",
    "photo_big": "фото: слишком большое",
}
EMAIL_REASONS = {
    "e_not_found": "emails.xlsx: карточка не найдена",
    "e_ambiguous": "emails.xlsx: неоднозначно",
    "e_alias": "emails.xlsx: алиас — карточка не найдена или неоднозначна",
    "e_bad": "emails.xlsx: почта некорректна",
    "e_dup": "emails.xlsx: почта уже занята (пропуск)",
}


class InputError(Exception):
    """Ошибка входных файлов; текст не содержит данных из файлов."""


# ---------- чтение xlsx ----------


def _norm(v) -> str:
    return " ".join(str(v or "").lower().replace("ё", "е").split())


def _cell(row, idx) -> str:
    v = row[idx] if idx is not None and idx < len(row) else None
    return "" if v is None else str(v).strip()


def read_sheet(path: Path, sheet_prefix: str, cols: dict[str, str]):
    """(строки [(номер строки Excel, значения)], {ключ: индекс столбца}). Столбцы ищутся по префиксу заголовка."""
    if not path.is_file():
        raise InputError(f"нет файла {path.name}")
    try:
        wb = load_workbook(path, data_only=True)
    except Exception as e:  # noqa: BLE001
        raise InputError(f"{path.name}: не удалось прочитать ({type(e).__name__})") from None
    try:
        ws = next((w for w in wb.worksheets if _norm(w.title).startswith(sheet_prefix)), None)
        if ws is None:
            raise InputError(f"{path.name}: нет листа «{sheet_prefix}…»")
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()
    if not rows:
        raise InputError(f"{path.name}: лист пуст")
    headers = [_norm(h) for h in rows[0]]
    index, missing = {}, []
    for key, prefix in cols.items():
        found = next((i for i, h in enumerate(headers) if h.startswith(prefix)), None)
        if found is None:
            missing.append(key)
        index[key] = found
    if missing:
        raise InputError(f"{path.name}: нет столбцов: {', '.join(missing)}")
    data = [(n, r) for n, r in enumerate(rows[1:], start=2) if any(v not in (None, "") and str(v).strip() for v in r)]
    return data, index


def read_aliases(path: Path | None) -> dict:
    if path is None:
        return {}
    try:
        raw = json.loads(path.read_text(encoding="utf-8-sig"))
    except Exception as e:  # noqa: BLE001
        raise InputError(f"aliases: не удалось прочитать JSON ({type(e).__name__})") from None
    if not isinstance(raw, dict) or not all(isinstance(k, str) and isinstance(v, str) for k, v in raw.items()):
        raise InputError("aliases: ожидается объект {\"Фамилия Имя\": \"Фамилия Имя\"}")
    return {service._name_words(k): service._name_words(v) for k, v in raw.items()}


# ---------- нормализация значений ----------


def norm_work_format(v: str) -> str:
    low = v.lower()
    if low.startswith("пятидневка"):
        return "Пятидневка в офисе"
    if "гибрид" in low:
        return "Гибрид"
    if "удал" in low:
        return "Удалёнка"
    return v


def norm_telegram(v: str) -> str | None:
    t = TG_PREFIX.sub("", v.strip()).lstrip("@").strip().rstrip("/")
    return t if TG_RE.match(t) else None


def _split_long(piece: str) -> list[str]:
    """Кусок >CHIP_MAX: по границам предложений, предложения упаковываются жадно до CHIP_MAX; слишком длинное предложение режется по пробелам."""
    out, cur = [], ""
    for sent in re.split(r"(?<=[.!?…])\s+", piece):
        if len(sent) > CHIP_MAX:
            if cur:
                out.append(cur)
                cur = ""
            while len(sent) > CHIP_MAX:
                cut = sent.rfind(" ", 0, CHIP_MAX + 1)
                cut = cut if cut > 0 else CHIP_MAX
                out.append(sent[:cut].strip())
                sent = sent[cut:].strip()
            cur = sent
        elif cur and len(cur) + 1 + len(sent) > CHIP_MAX:
            out.append(cur)
            cur = sent
        else:
            cur = f"{cur} {sent}".strip()
    if cur:
        out.append(cur)
    return [c for c in out if c]


def to_chips(v: str) -> list[str]:
    chips = []
    for piece in re.split(r"[\n;]+", v.replace("\r", "\n")):
        piece = piece.strip()
        if piece:
            chips.extend([piece] if len(piece) <= CHIP_MAX else _split_long(piece))
    return chips


def drive_id(v: str) -> str | None:
    m = DRIVE_ID.search(v) or DRIVE_PATH_ID.search(v)
    return m.group(1) if m else None


def read_photo(photos: Path, did: str) -> tuple[tuple[bytes, bytes] | None, str | None]:
    full_p, thumb_p = photos / f"{did}.full.jpg", photos / f"{did}.thumb.jpg"
    if not full_p.is_file() or not thumb_p.is_file():
        return None, "photo_missing"
    full, thumb = full_p.read_bytes(), thumb_p.read_bytes()
    if not full.startswith(service.JPEG_MAGIC) or not thumb.startswith(service.JPEG_MAGIC):
        return None, "photo_bad"
    if len(full) > service.FULL_LIMIT or len(thumb) > service.THUMB_LIMIT:
        return None, "photo_big"
    return (full, thumb), None


# ---------- транслитерация и нечёткое сопоставление почт ----------

TRANSLIT = {
    "а": "a", "б": "b", "в": "v", "г": "g", "д": "d", "е": "e", "ё": "e", "ж": "zh", "з": "z", "и": "i", "й": "y", "к": "k", "л": "l", "м": "m",
    "н": "n", "о": "o", "п": "p", "р": "r", "с": "s", "т": "t", "у": "u", "ф": "f", "х": "kh", "ц": "ts", "ч": "ch", "ш": "sh", "щ": "shch",
    "ъ": "", "ы": "y", "ь": "", "э": "e", "ю": "yu", "я": "ya",
}
# Варианты написания (ю: yu/iu/ju, я: ya/ia, х: kh/h, ц: ts/c, щ: sch/shch, и/й: i/y/j, ий: iy/y/i, ё: e/yo, ь/ъ: пусто) сводятся к одной форме:
# и карточка, и латинское имя из emails.xlsx приводятся к ней одинаково, затем сравниваются difflib.
FOLD = [("shch", "w"), ("sch", "w"), ("sh", "w"), ("zh", "z"), ("kh", "h"), ("ts", "c"), ("ch", "q"), ("yo", "e"), ("ye", "e"), ("yu", "u"),
        ("iu", "u"), ("ju", "u"), ("ya", "a"), ("ia", "a"), ("ja", "a"), ("x", "ks"), ("y", "i"), ("j", "i")]


def fold(word: str) -> str:
    s = re.sub(r"[^a-z]", "", word.lower())
    for a, b in FOLD:
        s = s.replace(a, b)
    return re.sub(r"(.)\1+", r"\1", s)


def fold_cyr(word: str) -> str:
    return fold("".join(TRANSLIT.get(c, c) for c in word.lower()))


_ratio_cache: dict = {}


def _ratio(a: str, b: str) -> float:
    key = (a, b)
    if key not in _ratio_cache:
        _ratio_cache[key] = SequenceMatcher(None, a, b).ratio()
    return _ratio_cache[key]


def assignment_score(row_words: list[str], card_words: list[str]) -> float | None:
    """Лучшее соответствие слов (порядок не важен): среднее, если каждое слово >= WORD_MIN; иначе None."""
    if len(row_words) != len(card_words) or not 1 <= len(row_words) <= MAX_WORDS:
        return None
    best = None
    for perm in itertools.permutations(card_words):
        scores = [_ratio(a, b) for a, b in zip(row_words, perm)]
        if min(scores) >= WORD_MIN:
            avg = sum(scores) / len(scores)
            best = avg if best is None or avg > best else best
    return best


def match_email_row(name: str, cards: list[dict]) -> tuple[dict | None, str | None]:
    row_words = [fold(w) for w in re.findall(r"[A-Za-z]+", name)]
    row_words = [w for w in row_words if w]
    scored = sorted(
        ((sc, c["id"], c) for c in cards if (sc := assignment_score(row_words, c["folded"])) is not None), key=lambda t: (-t[0], t[1])
    )
    if not scored:
        return None, "e_not_found"
    if len(scored) > 1 and scored[0][0] - scored[1][0] < LEAD_MIN:
        return None, "e_ambiguous"
    return scored[0][2], None


# ---------- импорт ----------


def _rows(nums, limit: int = 60) -> str:
    nums = sorted(nums)
    return ", ".join(map(str, nums[:limit])) + (", …" if len(nums) > limit else "")


def _empty(v) -> bool:
    return v in (None, "", [])


def valid_email(v: str) -> bool:
    return bool(service.EMAIL_RE.match(v)) and len(v) <= EMAIL_MAX


def answers_from_row(row, idx, issues: list) -> dict:
    """Нормализованные значения анкеты из строки ответа; проблемы — в issues (код причины)."""
    out = {}
    for key, limit in (("location", 80), ("work_format", 80), ("responsibility", 2000), ("ask_me_about", 2000)):
        v = _cell(row, idx[key])
        if key == "work_format" and v:
            v = norm_work_format(v)
        if not v:
            continue
        if len(v) > limit:
            issues.append("too_long")
        else:
            out[key] = v
    tg = _cell(row, idx["telegram"])
    if tg:
        t = norm_telegram(tg)
        if t is None:
            issues.append("telegram")
        else:
            out["telegram"] = t
    chips = to_chips(_cell(row, idx["hobbies"]))
    if len(chips) > CHIPS_MAX:
        chips = chips[:CHIPS_MAX]
        issues.append("chips_cut")
    if chips:
        out["hobbies"] = chips
    return out


def run_import(session: Session, folder: str | Path, aliases_path: str | Path | None, *, apply: bool) -> Report:
    head = f"Перенос ответов Google-формы в «Карту команды»: {'ЗАПИСЬ' if apply else 'СУХОЙ ПРОГОН (ничего не записано)'}"
    folder = Path(folder)
    try:
        if not folder.is_dir():
            raise InputError("папка не найдена")
        responses, ridx = read_sheet(folder / "responses.xlsx", "ответы на форму", RESPONSE_COLS)
        emails, eidx = read_sheet(folder / "emails.xlsx", "sheet1", EMAIL_COLS)
        aliases = read_aliases(Path(aliases_path) if aliases_path else None)
        if not (folder / "photos").is_dir():
            raise InputError("нет каталога photos/")
    except InputError as e:
        return Report(1, [head, f"Ошибка входных данных: {e}", "Итог: 1 — ничего не записано"])
    try:
        return _run(session, folder / "photos", responses, ridx, emails, eidx, aliases, apply, head)
    except DBAPIError as e:
        session.rollback()
        diag = getattr(getattr(e, "orig", None), "diag", None)  # сообщения psycopg содержат значения строк (ПДн)
        return Report(2, [head, f"Конфликт БД ({type(getattr(e, 'orig', e)).__name__}): таблица={getattr(diag, 'table_name', None)}, "
                                f"ограничение={getattr(diag, 'constraint_name', None)}. Ничего не записано.", "Итог: 2"])
    except BaseException:
        session.rollback()
        raise


def _run(session, photos: Path, responses, ridx, emails, eidx, aliases, apply: bool, head: str) -> Report:
    session.execute(text("SELECT pg_advisory_xact_lock(:k)"), {"k": LOCK_TEAM})
    rows = session.execute(
        text(f"""SELECT p.id, p.first_name, p.last_name, p.email, p.answers, ph.person_id IS NOT NULL AS has_photo
                 FROM {S}people p LEFT JOIN {S}photos ph ON ph.person_id = p.id WHERE NOT p.is_archived ORDER BY p.id""")
    ).mappings().all()
    used_emails = {e.lower() for e in session.execute(text(f"SELECT email FROM {S}people WHERE email IS NOT NULL")).scalars()}
    cards, by_words = {}, {}
    for r in rows:
        words = service._name_words(r["last_name"], r["first_name"])
        card = {
            "id": r["id"], "email": r["email"] or None, "old_answers": dict(r["answers"] or {}), "answers": dict(r["answers"] or {}), "new_answers": {},
            "new_email": None, "has_photo": r["has_photo"], "photo": None,
            "folded": [fold_cyr(w) for w in re.findall(r"\w+", f"{r['last_name']} {r['first_name']}".lower())],
        }
        cards[r["id"]] = card
        by_words.setdefault(words, []).append(card)

    bad: dict[str, list[int]] = {k: [] for k in {**REASONS, **EMAIL_REASONS}}
    n = {"matched": 0, "repeat": 0, "fields": 0, "photos": 0, "emails_resp": 0, "emails_xlsx": 0, "emails_has": 0}
    touched: set[int] = set()

    def find(words):
        found = by_words.get(words, [])
        if len(found) == 1:
            return found[0], None
        target = aliases.get(words)
        if target is not None:
            found = by_words.get(target, [])
            return (found[0], None) if len(found) == 1 else (None, "alias_not_found" if not found else "alias_ambiguous")
        return None, "not_found" if not found else "ambiguous"

    def set_email(card, value: str, counter: str, rownum: int, dup_code: str) -> bool:
        if value.lower() in used_emails:
            bad[dup_code].append(rownum)
            return False
        used_emails.add(value.lower())
        card["new_email"] = card["email"] = value
        touched.add(card["id"])
        n[counter] += 1
        return True

    seen_cards: set[int] = set()
    for rownum, row in reversed(responses):  # новые ответы раньше: при повторной анкете побеждает последняя
        words = service._name_words(_cell(row, ridx["last"]), _cell(row, ridx["first"]))
        if not words:
            bad["empty_name"].append(rownum)
            continue
        card, reason = find(words)
        if card is None:
            bad[reason].append(rownum)
            continue
        n["matched"] += 1
        n["repeat"] += card["id"] in seen_cards
        seen_cards.add(card["id"])
        issues: list[str] = []
        for key, v in answers_from_row(row, ridx, issues).items():
            if _empty(card["answers"].get(key)):
                card["answers"][key] = card["new_answers"][key] = v
                n["fields"] += 1
                touched.add(card["id"])
        em = _cell(row, ridx["email"])
        if em and not card["email"]:
            if valid_email(em):
                set_email(card, em, "emails_resp", rownum, "email_dup")
            else:
                issues.append("email_bad")
        if not card["has_photo"]:
            link = _cell(row, ridx["photo"])
            did = drive_id(link) if link else None
            if did is None:
                issues.append("photo_no_link")
            else:
                photo, code = read_photo(photos, did)
                if code:
                    issues.append(code)
                else:
                    card["photo"], card["has_photo"] = photo, True
                    n["photos"] += 1
                    touched.add(card["id"])
        for code in dict.fromkeys(issues):
            bad[code].append(rownum)

    for rownum, row in emails:
        em = _cell(row, eidx["email"])
        if not valid_email(em):
            bad["e_bad"].append(rownum)
            continue
        ename = _cell(row, eidx["name"])
        target = aliases.get(service._name_words(ename))  # алиас (в т.ч. латинский ключ) приоритетнее транслита
        if target is not None:
            found = by_words.get(target, [])
            card, code = (found[0], None) if len(found) == 1 else (None, "e_alias")
        else:
            card, code = match_email_row(ename, list(cards.values()))
        if card is None:
            bad[code].append(rownum)
        elif card["email"]:
            n["emails_has"] += 1
        else:
            set_email(card, em, "emails_xlsx", rownum, "e_dup")

    for cid in sorted(touched):
        card = cards[cid]
        cols = {"email": card["new_email"]} if card["new_email"] else {}
        if cols or card["new_answers"]:
            service._update_row(session, cid, cols, card["new_answers"], card["old_answers"])
        if card["photo"]:
            service._save_photo(session, cid, *card["photo"])

    lines = [
        head,
        f"Ответов в responses.xlsx: {len(responses)}, сопоставлено с карточками: {n['matched']} (повторных ответов на ту же карточку: {n['repeat']})",
        f"Строк в emails.xlsx: {len(emails)}, карточка уже с почтой: {n['emails_has']}",
        f"Карточек обновлено: {len(touched)}, полей анкеты записано: {n['fields']}, фото: {n['photos']}, "
        f"почт: {n['emails_resp'] + n['emails_xlsx']} (из ответов {n['emails_resp']}, из emails.xlsx {n['emails_xlsx']})",
    ]
    skipped = {**REASONS, **EMAIL_REASONS}
    total = sum(len(v) for v in bad.values())
    lines.append(f"Пропущено/замечаний: {total}")
    for code, label in skipped.items():
        if bad[code]:
            lines.append(f"  {label}: {len(bad[code])}; строки Excel: {_rows(bad[code])}")
    if apply:
        counters = {**n, "responses": len(responses), "email_rows": len(emails), "cards_updated": len(touched), "skipped": {k: len(v) for k, v in bad.items() if v}}
        log_action(session, None, ACTION, service.MODULE_ID, **counters)
        session.commit()
        return Report(0, lines + ["Итог: 0 — перенос записан"])
    session.rollback()
    return Report(0, lines + ["Итог: 0 — можно запускать с --apply"])


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(prog="python -m app.modules.team.import_responses", description="Перенос ответов Google-формы в карточки «Карты команды»")
    ap.add_argument("folder", help="папка с responses.xlsx, emails.xlsx и photos/")
    ap.add_argument("--aliases", help="aliases.json: {\"Фамилия Имя из ответа\": \"Фамилия Имя карточки\"}")
    ap.add_argument("--apply", action="store_true", help="записать (по умолчанию сухой прогон)")
    args = ap.parse_args(argv)
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    try:
        with Session(get_engine()) as session:
            report = run_import(session, args.folder, args.aliases, apply=args.apply)
    except Exception as e:  # noqa: BLE001 — текст исключения может содержать ПДн, печатаем только тип
        print(f"Непредвиденная ошибка: {type(e).__name__}", file=sys.stderr)
        return 1
    print(report.text())
    return report.exit_code


if __name__ == "__main__":
    raise SystemExit(main())
