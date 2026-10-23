"""Перенос ответов Google-формы (app.modules.team.import_responses). Все файлы синтетические, в tmp_path."""
import json

import pytest
from openpyxl import Workbook
from sqlalchemy import text

from app.modules.team import import_responses as ir
from app.modules.team import service

JPEG = b"\xff\xd8\xff\xe0" + b"\x00" * 50
HEADERS = [
    "Отметка времени", "Имя", "Фамилия ", "Твоя должность", "ТВОЙ РУКОВОДИТЕЛЬ", "Расскажи, какая у тебя сейчас зона ответственности?",
    "Приведи примеры, с чем к тебе можно обращаться", "В какой локации ты работаешь?", "Какой у тебя формат работы?", "Как тебя найти в телеграм?",
    "Какие еще у тебя есть таланты?", "Закинь сюда свою фотку", "Напиши свою корпоративную почту",
]


def resp_row(first, last, **kw):
    return [None, first, last, "Должность из формы", "Кто-то", kw.get("resp"), kw.get("ask"), kw.get("loc"), kw.get("fmt"), kw.get("tg"), kw.get("hob"),
            kw.get("photo"), kw.get("email")]


def make_folder(tmp_path, responses=(), emails=(), photos=None, sheet="Ответы на форму (1)"):
    ws_wb = Workbook()
    ws = ws_wb.active
    ws.title = sheet
    ws.append(HEADERS)
    for r in responses:
        ws.append(r)
    ws_wb.save(tmp_path / "responses.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Имя", "Почта"])
    for r in emails:
        ws.append(list(r))
    wb.save(tmp_path / "emails.xlsx")
    (tmp_path / "photos").mkdir(exist_ok=True)
    for did, (full, thumb) in (photos or {}).items():
        if full is not None:
            (tmp_path / "photos" / f"{did}.full.jpg").write_bytes(full)
        if thumb is not None:
            (tmp_path / "photos" / f"{did}.thumb.jpg").write_bytes(thumb)
    return tmp_path


def add_person(db, first, last, email=None, answers=None, archived=False, tag="тег", team="команда"):
    return db.execute(
        text(
            """INSERT INTO trendsetter_hub.team_people (first_name, last_name, title, direction_id, email, answers, is_archived, tag, team)
               VALUES (:f, :l, 'Должность', 'ops', :e, CAST(:a AS jsonb), :ar, :tag, :team) RETURNING id"""
        ),
        {"f": first, "l": last, "e": email, "a": json.dumps(answers or {}, ensure_ascii=False), "ar": archived, "tag": tag, "team": team},
    ).scalar_one()


def person(db, pid):
    db.rollback()
    return db.execute(text("SELECT * FROM trendsetter_hub.team_people WHERE id = :i"), {"i": pid}).mappings().one()


def has_photo(db, pid):
    db.rollback()
    return db.execute(text("SELECT count(*) FROM trendsetter_hub.team_photos WHERE person_id = :i"), {"i": pid}).scalar_one() == 1


def run(db, folder, apply=True, aliases=None):
    return ir.run_import(db, folder, aliases, apply=apply)


# ---------- сопоставление ----------


def test_match_any_order_yo_and_only_unarchived(db, tmp_path):
    a = add_person(db, "Пётр", "Фёдоров")
    gone = add_person(db, "Анна", "Архивова", archived=True)
    db.commit()
    folder = make_folder(tmp_path, [resp_row("Петр", "ФЕДОРОВ", loc="Москва"), resp_row("Анна", "Архивова", loc="Москва")])
    r = run(db, folder)
    assert r.exit_code == 0, r.text()
    assert person(db, a)["answers"] == {"location": "Москва"}
    assert person(db, gone)["answers"] == {}
    assert "сопоставлено с карточками: 1" in r.text() and "нет подходящей карточки: 1; строки Excel: 3" in r.text()


def test_ambiguous_and_not_found_reported_by_row_not_name(db, tmp_path):
    add_person(db, "Иван", "Двойников")
    add_person(db, "Иван", "Двойников")
    db.commit()
    folder = make_folder(tmp_path, [resp_row("Иван", "Двойников", loc="Москва"), resp_row("Никто", "Неизвестный", loc="Москва")])
    r = run(db, folder)
    out = r.text()
    assert "несколько подходящих карточек (нужен алиас): 1; строки Excel: 2" in out and "нет подходящей карточки: 1; строки Excel: 3" in out
    assert "Двойников" not in out and "Неизвестный" not in out
    assert db.execute(text("SELECT count(*) FROM trendsetter_hub.team_people WHERE answers <> '{}'::jsonb")).scalar_one() == 0


def test_aliases(db, tmp_path):
    pid = add_person(db, "Елизавета", "Зимина")
    db.commit()
    folder = make_folder(tmp_path, [resp_row("Лиза", "Зимина", loc="Ереван"), resp_row("Маша", "Левая", loc="Ереван")])
    (tmp_path / "aliases.json").write_text(json.dumps({"зимина лиза": "Зимина  Елизавета", "Левая Маша": "Нет Такой"}, ensure_ascii=False), encoding="utf-8")
    r = run(db, folder, aliases=tmp_path / "aliases.json")
    assert person(db, pid)["answers"] == {"location": "Ереван"}
    assert "алиас: карточка не найдена: 1; строки Excel: 3" in r.text()


# ---------- только пустое ----------


def test_fills_only_empty_and_never_touches_org_fields(db, tmp_path):
    pid = add_person(db, "Иван", "Иванов", email="old@example.com", answers={"location": "Ереван", "hobbies": ["чай"], "telegram": ""}, tag="СТМ", team="Т1")
    db.execute(text("UPDATE trendsetter_hub.team_people SET show_in_map = false WHERE id = :i"), {"i": pid})
    db.commit()
    folder = make_folder(
        tmp_path, [resp_row("Иван", "Иванов", loc="Москва", fmt="Гибридный формат", tg="@ivan_1", hob="кофе", resp="Зона", email="new@example.com")]
    )
    run(db, folder)
    row = person(db, pid)
    assert row["answers"] == {"location": "Ереван", "hobbies": ["чай"], "telegram": "ivan_1", "work_format": "Гибрид", "responsibility": "Зона"}
    assert row["email"] == "old@example.com" and row["title"] == "Должность" and row["tag"] == "СТМ" and row["team"] == "Т1"
    assert row["show_in_map"] is False and row["manager_id"] is None and row["direction_id"] == "ops" and row["user_id"] is None
    assert (row["first_name"], row["last_name"]) == ("Иван", "Иванов")


def test_email_from_response_then_emails_xlsx_and_duplicates(db, tmp_path):
    a = add_person(db, "Иван", "Иванов")
    b = add_person(db, "Мария", "Петрова")
    c = add_person(db, "Олег", "Сидоров")
    add_person(db, "Занят", "Почтов", email="Busy@Example.com")
    db.commit()
    folder = make_folder(
        tmp_path,
        [resp_row("Иван", "Иванов", email="ivan@example.com"), resp_row("Олег", "Сидоров", email="busy@example.com")],
        emails=[("Maria Petrova", "maria@example.com"), ("Ivan Ivanov", "other@example.com"), ("Oleg Sidorov", "oleg@example.com"), ("Nobody Here", "n@example.com")],
    )
    r = run(db, folder)
    assert person(db, a)["email"] == "ivan@example.com"  # из ответа, emails.xlsx не перезаписывает
    assert person(db, b)["email"] == "maria@example.com"
    assert person(db, c)["email"] == "oleg@example.com"  # дубль в ответе пропущен, взято из emails.xlsx
    out = r.text()
    assert "почта уже занята другой карточкой (пропуск): 1; строки Excel: 3" in out
    assert "emails.xlsx: карточка не найдена: 1; строки Excel: 5" in out and "карточка уже с почтой: 1" in out
    assert "почт: 3 (из ответов 1, из emails.xlsx 2)" in out


def test_email_xlsx_duplicate_value_and_invalid(db, tmp_path):
    a = add_person(db, "Иван", "Иванов")
    b = add_person(db, "Мария", "Петрова")
    db.commit()
    folder = make_folder(tmp_path, emails=[("Ivan Ivanov", "same@example.com"), ("Maria Petrova", "SAME@example.com"), ("Maria Petrova", "not-an-email")])
    r = run(db, folder)
    assert person(db, a)["email"] == "same@example.com" and person(db, b)["email"] is None
    assert "emails.xlsx: почта уже занята (пропуск): 1; строки Excel: 3" in r.text() and "emails.xlsx: почта некорректна: 1; строки Excel: 4" in r.text()


@pytest.mark.parametrize(
    "ru_first,ru_last,lat",
    [
        ("Юлия", "Хмелёва", "Iuliia Khmeleva"), ("Алексей", "Серов", "Aleksey Serov"), ("Алексей", "Серов", "Serov Aleksey"), ("Дарья", "Щукина", "Shchukina Darya"),
        ("Юрий", "Цой", "Yuri Tsoy"), ("Евгений", "Чехов", "Evgeny Chekhov"), ("Мария", "Жуковская", "Maria Zhukovskaya"),
    ],
)
def test_translit_variants(db, tmp_path, ru_first, ru_last, lat):
    pid = add_person(db, ru_first, ru_last)
    add_person(db, "Совсем", "Другой")
    db.commit()
    run(db, make_folder(tmp_path, emails=[(lat, "x@example.com")]))
    assert person(db, pid)["email"] == "x@example.com"


def test_email_ambiguous_close_candidates(db, tmp_path):
    add_person(db, "Алексей", "Серов")
    add_person(db, "Алексей", "Серова")
    db.commit()
    r = run(db, make_folder(tmp_path, emails=[("Aleksey Serov", "x@example.com")]))
    assert "emails.xlsx: неоднозначно: 1; строки Excel: 2" in r.text()
    assert db.execute(text("SELECT count(*) FROM trendsetter_hub.team_people WHERE email IS NOT NULL")).scalar_one() == 0


# ---------- нормализация ----------


@pytest.mark.parametrize(
    "raw,expected",
    [("@user_1", "user_1"), ("https://t.me/user_1", "user_1"), ("t.me/user_1", "user_1"), ("telegram.me/user_1", "user_1"), ("user_1", "user_1"), ("@ab", None), ("user name", None), ("@иван", None)],
)
def test_telegram(raw, expected):
    assert ir.norm_telegram(raw) == expected


@pytest.mark.parametrize(
    "raw,expected",
    [("Пятидневка в офисе, очень", "Пятидневка в офисе"), ("пятидневка", "Пятидневка в офисе"), ("Гибридный", "Гибрид"), ("Полностью удалённо", "Удалёнка"), ("Сменный график", "Сменный график")],
)
def test_work_format(raw, expected):
    assert ir.norm_work_format(raw) == expected


def test_chips_splitting():
    assert ir.to_chips("чай; кофе\r\n\n  ; йога ") == ["чай", "кофе", "йога"]
    sent = "Я люблю " + "очень " * 30 + "долго."  # ~200 символов
    long = f"{sent} {sent} {sent}"
    chips = ir.to_chips(long)
    assert len(chips) == 2 and all(len(c) <= 400 for c in chips) and " ".join(chips) == long
    assert all(len(c) <= 400 for c in ir.to_chips("слово " * 200))  # одно предложение без точек — режется по пробелам


def test_chips_limit_20_and_long_values_skipped(db, tmp_path):
    pid = add_person(db, "Иван", "Иванов")
    db.commit()
    folder = make_folder(tmp_path, [resp_row("Иван", "Иванов", hob="; ".join(f"т{i}" for i in range(25)), resp="я" * 2001, loc="м" * 81, ask="я" * 2000)])
    r = run(db, folder)
    a = person(db, pid)["answers"]
    assert len(a["hobbies"]) == 20 and "responsibility" not in a and "location" not in a and len(a["ask_me_about"]) == 2000
    assert "чипов больше 20 (лишние отброшены): 1; строки Excel: 2" in r.text()
    assert "слишком длинное (поле пропущено): 1; строки Excel: 2" in r.text()


def test_imported_chip_passes_service_validation():
    field = {"kind": "chips", "key": "hobbies", "label": "Увлечения"}
    assert service._clean_value(field, ["я" * 400]) == ["я" * 400]
    with pytest.raises(Exception):
        service._clean_value(field, ["я" * 401])


# ---------- фото ----------


def test_photos(db, tmp_path):
    ok = add_person(db, "Иван", "Иванов")
    existing = add_person(db, "Мария", "Петрова")
    add_person(db, "Олег", "Сидоров")
    add_person(db, "Нина", "Нинова")
    add_person(db, "Вера", "Верова")
    add_person(db, "Зоя", "Зоева")
    db.execute(text("INSERT INTO trendsetter_hub.team_photos (person_id, photo_full, photo_thumb) VALUES (:i, :f, :t)"), {"i": existing, "f": b"\xff\xd8\xffOLD", "t": b"\xff\xd8\xffOLD"})
    db.commit()
    link = lambda d: f"https://drive.google.com/open?id={d}"  # noqa: E731
    folder = make_folder(
        tmp_path,
        [
            resp_row("Иван", "Иванов", photo=link("OK1")), resp_row("Мария", "Петрова", photo=link("OK1")),
            resp_row("Олег", "Сидоров", photo=link("NOFILE")), resp_row("Нина", "Нинова", photo=link("BAD")),
            resp_row("Вера", "Верова", photo=link("BIG")), resp_row("Зоя", "Зоева", photo=None),
        ],
        photos={"OK1": (JPEG, JPEG), "BAD": (b"GIF89a" + b"\x00" * 10, JPEG), "BIG": (JPEG + b"\x00" * service.FULL_LIMIT, JPEG), "NOFILE": (JPEG, None)},
    )
    r = run(db, folder)
    assert has_photo(db, ok)
    assert bytes(db.execute(text("SELECT photo_full FROM trendsetter_hub.team_photos WHERE person_id = :i"), {"i": existing}).scalar_one()) == b"\xff\xd8\xffOLD"
    out = r.text()
    for label, row in (("нет файла full/thumb", 4), ("не JPEG", 5), ("слишком большое", 6), ("нет ссылки/ID", 7)):
        assert f"{label}: 1; строки Excel: {row}" in out
    assert "фото: 1" in out
    assert db.execute(text("SELECT count(*) FROM trendsetter_hub.team_photos")).scalar_one() == 2


# ---------- режимы ----------


def test_dry_run_apply_and_repeat(db, tmp_path):
    pid = add_person(db, "Иван", "Иванов")
    db.commit()
    folder = make_folder(
        tmp_path, [resp_row("Иван", "Иванов", loc="Москва", tg="@ivan_1", photo="https://drive.google.com/open?id=P1", email="ivan@example.com")], photos={"P1": (JPEG, JPEG)}
    )
    dry = run(db, folder, apply=False)
    assert dry.exit_code == 0 and "СУХОЙ ПРОГОН" in dry.text() and "Карточек обновлено: 1, полей анкеты записано: 2, фото: 1, почт: 1" in dry.text()
    assert person(db, pid)["answers"] == {} and person(db, pid)["email"] is None and not has_photo(db, pid)
    assert db.execute(text("SELECT count(*) FROM trendsetter_hub.audit_log WHERE action = 'team.import_responses'")).scalar_one() == 0
    done = run(db, folder, apply=True)
    assert done.exit_code == 0
    assert person(db, pid)["answers"] == {"location": "Москва", "telegram": "ivan_1"} and person(db, pid)["email"] == "ivan@example.com" and has_photo(db, pid)
    again = run(db, folder, apply=True)
    assert again.exit_code == 0 and "Карточек обновлено: 0, полей анкеты записано: 0, фото: 0, почт: 0" in again.text()
    audit = db.execute(text("SELECT details FROM trendsetter_hub.audit_log WHERE action = 'team.import_responses' ORDER BY id")).scalars().all()
    assert len(audit) == 2 and audit[0]["cards_updated"] == 1 and audit[0]["fields"] == 2


def test_no_personal_data_in_output_and_audit(db, tmp_path):
    add_person(db, "Секретный", "Человеков")
    db.commit()
    folder = make_folder(
        tmp_path,
        [resp_row("Секретный", "Человеков", loc="Тайгород", tg="secret_tg", hob="тайное хобби", resp="Секретная зона", email="secret@example.com"), resp_row("Чужой", "Человек", tg="@x")],
        emails=[("Sekretnyy Chelovekov", "other@example.com"), ("Unknown Person", "unknown@example.com")],
    )
    for apply in (False, True):
        out = run(db, folder, apply=apply).text()
        for secret in ("Секретный", "Человеков", "Тайгород", "secret_tg", "хобби", "Секретная", "secret@example.com", "unknown@example.com", "Unknown", "Чужой"):
            assert secret not in out
    audit = db.execute(text("SELECT details::text FROM trendsetter_hub.audit_log WHERE action = 'team.import_responses'")).scalar_one()
    for secret in ("Секретный", "Человеков", "Тайгород", "secret_tg", "secret@example.com"):
        assert secret not in audit


# ---------- входные ошибки ----------


def test_input_errors_exit_1(db, tmp_path):
    assert run(db, tmp_path / "nope").exit_code == 1
    assert run(db, tmp_path).exit_code == 1  # нет responses.xlsx
    (tmp_path / "ok").mkdir()
    folder = make_folder(tmp_path / "ok")
    assert run(db, folder).exit_code == 0
    wb = Workbook()
    wb.active.title = "Ответы на форму"
    wb.active.append(["Имя", "Фамилия"])
    wb.save(folder / "responses.xlsx")
    r = run(db, folder)
    assert r.exit_code == 1 and "нет столбцов" in r.text()
    make_folder(folder, sheet="Другой лист")
    assert run(db, folder).exit_code == 1
    make_folder(folder)
    (folder / "bad.json").write_text("{oops", encoding="utf-8")
    assert run(db, folder, aliases=folder / "bad.json").exit_code == 1


def test_main_exit_codes(tmp_path, capsys):
    assert ir.main([str(tmp_path / "missing")]) == 1  # папка проверяется до обращения к данным


def test_aliases_apply_to_emails_xlsx(db, tmp_path):
    a = add_person(db, "Софья", "Воробьёва")
    add_person(db, "Соня", "Воробьева")  # транслит выбрал бы неоднозначно/другую карточку; алиас приоритетнее
    db.commit()
    folder = make_folder(tmp_path, emails=[("Sonya Vorobyeva", "sonya@example.com"), ("Ghost Person", "ghost@example.com")])
    (tmp_path / "aliases.json").write_text(json.dumps({"sonya vorobyeva": "Воробьёва Софья", "Ghost Person": "Нет Такой"}), encoding="utf-8")
    r = run(db, folder, aliases=tmp_path / "aliases.json")
    assert person(db, a)["email"] == "sonya@example.com"
    assert "алиас — карточка не найдена или неоднозначна: 1; строки Excel: 3" in r.text()
