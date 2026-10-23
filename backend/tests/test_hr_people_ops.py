"""HR: напоминания и дни рождения, экспорт, история, чек-листы, оборудование, график отпусков, свои документы. Люди и данные вымышленные."""
import base64
import io
from datetime import date, timedelta

import pytest
from openpyxl import load_workbook
from sqlalchemy import text

from tests.conftest import login, make_user
from tests.test_hr import BASE, TODAY, PDF, admin, audit, b64, doc_body, grant, hr, iso, link, notes, person, resp, team, upload  # noqa: F401
from tests.test_team import add_person

TEAM = "/modules/team"
SALARY = 987654.32
SECRET_PASSPORT = "AB 1234567 секретный"
SECRET_DOC_TITLE = "Секретное соглашение"


def tpl_id(client, headers, kind="onboarding") -> int:
    return next(t["id"] for t in client.get(f"{BASE}/checklist-templates", headers=headers).json() if t["kind"] == kind)


@pytest.fixture()
def viewer(client, db):
    make_user(db, "v@example.com", "viewer")
    return login(client, "v@example.com")


@pytest.fixture()
def outsider(client, db):
    make_user(db, "o@example.com", "employee", name="Посторонний Пётр")
    return login(client, "o@example.com")


# ---------- права ----------


def test_hr_endpoints_rights_matrix(client, db, hr, resp, team, viewer, outsider, person):
    pid = person
    requests_ = [
        ("get", f"/people/{pid}/history", None),
        ("post", f"/people/{pid}/history", {"kind": "other", "effectiveDate": iso(0), "title": "Запись"}),
        ("get", f"/people/{pid}/equipment", None),
        ("post", f"/people/{pid}/equipment", {"kind": "laptop", "name": "Ноутбук", "issuedAt": iso(0)}),
        ("get", f"/people/{pid}/checklists", None),
        ("get", "/checklist-templates", None),
        ("get", "/export/people", None),
        ("get", f"/export/leave?from={iso(-30)}&to={iso(30)}", None),
        ("get", f"/leave/plan?year={TODAY.year}", None),
    ]
    for method, path, body in requests_:
        kw = {"json": body} if body is not None else {}
        assert getattr(client, method)(BASE + path, **kw).status_code == 401, path
        for who in (viewer, outsider, team["e"], team["m"]):
            assert getattr(client, method)(BASE + path, headers=who, **kw).status_code == 403, (path, who)
        for who in (hr, resp):
            assert getattr(client, method)(BASE + path, headers=who, **kw).status_code == 200, (path, who)


def test_me_endpoints_open_to_any_active_user(client, outsider, viewer):
    for who in (outsider, viewer):
        assert client.get(f"{BASE}/me/documents", headers=who).json() == []
        assert client.get(f"{BASE}/checklists/mine", headers=who).json() == []
        assert client.get(f"{BASE}/me/leave/plans", headers=who).json() == []
    assert client.get(f"{BASE}/me/documents").status_code == 401


def test_document_flag_and_patch_rights(client, hr, resp, person):
    r = upload(client, resp, person, visibleToEmployee=True)
    assert r.json()["visibleToEmployee"] is True
    assert upload(client, resp, person).json()["visibleToEmployee"] is False  # по умолчанию скрыт
    doc = r.json()["id"]
    assert client.patch(f"{BASE}/documents/{doc}", json={"visibleToEmployee": False}, headers=hr).status_code == 403
    done = client.patch(f"{BASE}/documents/{doc}", json={"visibleToEmployee": False, "title": "Новое"}, headers=resp)
    assert done.json()["visibleToEmployee"] is False and done.json()["title"] == "Новое"
    assert client.patch(f"{BASE}/documents/{doc}", json={"category": "nope"}, headers=resp).status_code == 422
    assert client.patch(f"{BASE}/documents/9999", json={"title": "x"}, headers=resp).status_code == 404
    assert client.get(f"{BASE}/people/{person}/documents", headers=resp).json()[1]["visibleToEmployee"] is False


# ---------- свои документы ----------


def test_my_documents_only_visible_and_own(client, team, resp, outsider, viewer):
    e, c = team["ids"]["e"], team["ids"]["c"]
    shown = upload(client, resp, e, title="Для меня", visibleToEmployee=True).json()
    hidden = upload(client, resp, e, title="Скрытый").json()
    other = upload(client, resp, c, title="Чужой", visibleToEmployee=True).json()
    mine = client.get(f"{BASE}/me/documents", headers=team["e"]).json()
    assert [d["id"] for d in mine] == [shown["id"]]
    assert set(mine[0]) == {"id", "category", "title", "fileName", "sizeBytes", "uploadedAt"}
    got = client.get(f"{BASE}/me/documents/{shown['id']}", headers=team["e"]).json()
    assert base64.b64decode(got["dataBase64"]) == PDF and got["mime"] == "application/pdf"
    assert client.get(f"{BASE}/me/documents/{hidden['id']}", headers=team["e"]).status_code == 404
    assert client.get(f"{BASE}/me/documents/{other['id']}", headers=team["e"]).status_code == 404
    assert client.get(f"{BASE}/me/documents/{shown['id']}", headers=team["c"]).status_code == 404
    assert client.get(f"{BASE}/me/documents/{shown['id']}", headers=outsider).status_code == 404
    assert client.get(f"{BASE}/me/documents/{shown['id']}", headers=viewer).status_code == 404
    # флаг снят — документ пропадает
    client.patch(f"{BASE}/documents/{shown['id']}", json={"visibleToEmployee": False}, headers=resp)
    assert client.get(f"{BASE}/me/documents", headers=team["e"]).json() == []
    assert client.get(f"{BASE}/me/documents/{shown['id']}", headers=team["e"]).status_code == 404


def test_my_document_download_is_logged_without_names(client, db, team, resp):
    shown = upload(client, resp, team["ids"]["e"], visibleToEmployee=True, fileName="паспорт-иванова.pdf").json()
    client.get(f"{BASE}/me/documents/{shown['id']}", headers=team["e"])
    details = audit(db, "hr.document_download")
    assert len(details) == 1 and "паспорт" not in str(details)


def test_my_documents_only_by_user_id_not_name_or_email(client, db, resp):
    """Совпадение имени или почты с чужой непривязанной карточкой документов не открывает; после привязки видны только отмеченные."""
    card = add_person(db, "Ирина", "Иванова", email="irina@example.com")
    doc = upload(client, resp, card, title="Мой", visibleToEmployee=True).json()
    upload(client, resp, card, title="Скрытый")
    make_user(db, "irina@example.com", "employee", name="Иванова Ирина")
    h = login(client, "irina@example.com")
    assert client.get(f"{BASE}/me/documents", headers=h).json() == []
    assert client.get(f"{BASE}/me/documents/{doc['id']}", headers=h).status_code == 404
    user_id = db.execute(text("SELECT id FROM trendsetter_hub.users WHERE email = 'irina@example.com'")).scalar_one()
    db.execute(text("UPDATE trendsetter_hub.team_people SET user_id = :u WHERE id = :p"), {"u": user_id, "p": card})
    db.commit()
    assert [d["title"] for d in client.get(f"{BASE}/me/documents", headers=h).json()] == ["Мой"]
    assert client.get(f"{BASE}/me/documents/{doc['id']}", headers=h).status_code == 200


# ---------- дни рождения ----------


def set_birth(client, hr, pid, value, **extra):
    r = client.patch(f"{BASE}/people/{pid}/file", json={"birthDate": value, **extra}, headers=hr)
    assert r.status_code == 200, r.text


def test_birthdays_day_and_month_only(client, db, hr, team, viewer):
    e, c, m = team["ids"]["e"], team["ids"]["c"], team["ids"]["m"]
    set_birth(client, hr, e, "1990-05-17")
    set_birth(client, hr, c, "1985-05-03", birthdayHidden=True)
    set_birth(client, hr, m, "1979-06-30")
    got = client.get(f"{BASE}/birthdays?month=2026-05", headers=team["e"])
    assert got.status_code == 200
    assert got.json() == [{"personId": e, "name": "Ермолов Егор", "directionId": "ops", "day": 17, "month": 5}]
    assert "1990" not in got.text and "1985" not in got.text
    assert [b["personId"] for b in client.get(f"{BASE}/birthdays?month=2026-06", headers=team["c"]).json()] == [m]
    assert client.get(f"{BASE}/people/{c}/file", headers=hr).json()["file"]["birthdayHidden"] is True
    # архивные и скрытые на карте не показываются
    db.execute(text("UPDATE trendsetter_hub.team_people SET is_archived = true WHERE id = :i"), {"i": e})
    db.commit()
    assert client.get(f"{BASE}/birthdays?month=2026-05", headers=team["m"]).json() == []
    assert client.get(f"{BASE}/birthdays?month=2026-13", headers=team["m"]).status_code == 422
    assert client.get(f"{BASE}/birthdays?month=bad", headers=team["m"]).status_code == 422
    assert client.get(f"{BASE}/birthdays", headers=team["m"]).status_code == 200  # текущий месяц по умолчанию
    assert client.get(f"{BASE}/birthdays", headers=viewer).status_code == 403  # viewer видит только схему
    assert client.patch(f"{BASE}/people/{c}/file", json={"birthdayHidden": "maybe"}, headers=hr).status_code == 422


def test_birthday_hidden_toggle_is_hr_only(client, hr, team):
    client.patch(f"{BASE}/people/{team['ids']['e']}/file", json={"hiredAt": "2024-01-01"}, headers=hr)
    assert client.patch(f"{BASE}/people/{team['ids']['e']}/file", json={"birthdayHidden": True}, headers=team["e"]).status_code == 403
    assert "birthdayHidden" not in client.get(f"{BASE}/me/file", headers=team["e"]).json()["file"]


# ---------- история ----------


def history_body(**over):
    return {"kind": "other", "effectiveDate": iso(-3), "title": "Перевод в группу", "oldValue": "А", "newValue": "Б", "comment": "ок", **over}


def test_history_salary_visible_only_to_responsible(client, db, hr, resp, person):
    assert client.post(f"{BASE}/people/{person}/history", json=history_body(kind="salary", title="Повышение оклада"), headers=hr).status_code == 403
    manual = client.post(f"{BASE}/people/{person}/history", json=history_body(), headers=hr).json()
    salary = client.post(f"{BASE}/people/{person}/history", json=history_body(kind="salary", title="Повышение оклада", newValue="150000"), headers=resp).json()
    assert salary["kind"] == "salary" and manual["auto"] is False and manual["createdBy"] == "Эйч Арова"
    assert [h["id"] for h in client.get(f"{BASE}/people/{person}/history", headers=hr).json()] == [manual["id"]]
    assert {h["id"] for h in client.get(f"{BASE}/people/{person}/history", headers=resp).json()} == {manual["id"], salary["id"]}
    assert "150000" not in client.get(f"{BASE}/people/{person}/history", headers=hr).text
    # удаление: HR не видит и не удаляет оклад
    assert client.post(f"{BASE}/history/{salary['id']}/delete", headers=hr).status_code == 404
    assert client.post(f"{BASE}/history/{salary['id']}/delete", headers=resp).json() == {"ok": True}
    assert client.post(f"{BASE}/history/{manual['id']}/delete", headers=hr).json() == {"ok": True}
    assert client.post(f"{BASE}/history/{manual['id']}/delete", headers=hr).status_code == 404
    # в журнале нет заголовков и значений
    assert all(set(d) == {"history_id", "person_id", "kind"} for d in audit(db, "hr.history_add") + audit(db, "hr.history_delete"))


def test_history_validation(client, hr, person):
    url = f"{BASE}/people/{person}/history"
    assert client.post(url, json=history_body(kind="nope"), headers=hr).status_code == 422
    assert client.post(url, json=history_body(title=" "), headers=hr).status_code == 422
    assert client.post(url, json=history_body(title="я" * 201), headers=hr).status_code == 422
    assert client.post(url, json=history_body(effectiveDate="31.12.2025"), headers=hr).status_code == 422
    assert client.post(f"{BASE}/people/9999/history", json=history_body(), headers=hr).status_code == 404


def test_history_auto_written_on_person_patch(client, db, hr, resp, team):
    e, c = team["ids"]["e"], team["ids"]["c"]
    url = f"{TEAM}/hr/people/{e}"
    assert client.patch(url, json={"title": "Должность", "directionId": "ops"}, headers=hr).status_code == 200  # без изменений
    assert client.get(f"{BASE}/people/{e}/history", headers=hr).json() == []
    assert client.patch(url, json={"title": "Старший инженер", "managerId": c, "directionId": "mkt"}, headers=hr).status_code == 200
    items = {h["kind"]: h for h in client.get(f"{BASE}/people/{e}/history", headers=hr).json()}
    assert set(items) == {"position", "manager", "direction"} and all(h["auto"] for h in items.values())
    assert items["position"]["oldValue"] == "Должность" and items["position"]["newValue"] == "Старший инженер"
    assert items["manager"]["oldValue"] == "Мэнэджер Мария" and items["manager"]["newValue"] == "Коллегин Клим"
    assert items["direction"]["newValue"] == "МАРКЕТИНГ" and items["direction"]["effectiveDate"] == iso(0)
    assert client.patch(url, json={"isArchived": True}, headers=hr).status_code == 200
    status = [h for h in client.get(f"{BASE}/people/{e}/history", headers=hr).json() if h["kind"] == "status"]
    assert status[0]["oldValue"] == "Работает" and status[0]["newValue"] == "Уволен"
    # автозапись нельзя удалить без responsible
    assert client.post(f"{BASE}/history/{status[0]['id']}/delete", headers=hr).status_code == 403
    assert client.post(f"{BASE}/history/{status[0]['id']}/delete", headers=resp).status_code == 200


# ---------- оборудование ----------


def test_equipment_crud_and_limit(client, db, hr, person):
    url = f"{BASE}/people/{person}/equipment"
    item = client.post(url, json={"kind": "laptop", "name": "MacBook Pro", "serial": "C02XYZ", "issuedAt": iso(-5), "returnDue": iso(30)}, headers=hr).json()
    assert item["returnedAt"] is None and item["serial"] == "C02XYZ"
    done = client.patch(f"{BASE}/equipment/{item['id']}", json={"returnedAt": iso(0), "comment": "вернул"}, headers=hr).json()
    assert done["returnedAt"] == iso(0) and done["comment"] == "вернул" and done["name"] == "MacBook Pro"
    assert client.patch(f"{BASE}/equipment/{item['id']}", json={"name": None}, headers=hr).status_code == 422
    assert client.patch(f"{BASE}/equipment/{item['id']}", json={"kind": "car"}, headers=hr).status_code == 422
    assert client.patch(f"{BASE}/equipment/{item['id']}", json={"bogus": 1}, headers=hr).status_code == 422
    assert client.patch(f"{BASE}/equipment/9999", json={"comment": "x"}, headers=hr).status_code == 404
    assert client.post(url, json={"kind": "laptop", "name": " ", "issuedAt": iso(0)}, headers=hr).status_code == 422
    assert client.post(url, json={"kind": "laptop", "name": "X", "issuedAt": "bad"}, headers=hr).status_code == 422
    assert len(client.get(url, headers=hr).json()) == 1
    assert all("MacBook" not in str(d) and "C02XYZ" not in str(d) for a in ("hr.equipment_add", "hr.equipment_update") for d in audit(db, a))
    db.execute(
        text("INSERT INTO trendsetter_hub.hr_equipment (person_id, kind, name, issued_at) SELECT :p, 'other', 'н', current_date FROM generate_series(1, 99)"),
        {"p": person},
    )
    db.commit()
    assert client.post(url, json={"kind": "other", "name": "лишнее", "issuedAt": iso(0)}, headers=hr).status_code == 409
    assert client.post(f"{BASE}/equipment/{item['id']}/delete", headers=hr).json() == {"ok": True}
    assert client.post(url, json={"kind": "other", "name": "теперь можно", "issuedAt": iso(0)}, headers=hr).status_code == 200


# ---------- чек-листы ----------


def test_templates_seeded_and_editable(client, hr):
    templates = client.get(f"{BASE}/checklist-templates", headers=hr).json()
    assert {t["kind"] for t in templates} == {"onboarding", "offboarding"} and all(5 <= len(t["items"]) <= 7 for t in templates)
    tid = templates[0]["id"]
    done = client.patch(f"{BASE}/checklist-templates/{tid}", json={"title": "Новое имя", "items": [{"title": "Один", "daysOffset": 2}, {"title": "Два"}]}, headers=hr).json()
    assert done["title"] == "Новое имя" and [(i["title"], i["daysOffset"], i["sort"]) for i in done["items"]] == [("Один", 2, 1), ("Два", None, 2)]
    assert client.patch(f"{BASE}/checklist-templates/{tid}", json={"active": False}, headers=hr).json()["active"] is False
    assert client.patch(f"{BASE}/checklist-templates/{tid}", json={"items": []}, headers=hr).status_code == 422
    assert client.patch(f"{BASE}/checklist-templates/9999", json={"active": True}, headers=hr).status_code == 404
    created = client.post(f"{BASE}/checklist-templates", json={"kind": "onboarding", "title": "Свой", "items": [{"title": "П"}]}, headers=hr)
    assert created.status_code == 200 and created.json()["active"] is True
    assert client.post(f"{BASE}/checklist-templates", json={"kind": "x", "title": "Свой", "items": [{"title": "П"}]}, headers=hr).status_code == 422


def test_checklist_flow_assignee_and_completion(client, db, hr, resp, team):
    e = team["ids"]["e"]
    new = client.patch(f"{BASE}/checklist-templates/{tpl_id(client, hr)}", json={"items": [{"title": "Аккаунты", "daysOffset": 1}, {"title": "Наставник"}]}, headers=hr).json()
    assert len(new["items"]) == 2
    inactive = client.post(f"{BASE}/checklist-templates", json={"kind": "onboarding", "title": "Выкл", "items": [{"title": "П"}]}, headers=hr).json()
    client.patch(f"{BASE}/checklist-templates/{inactive['id']}", json={"active": False}, headers=hr)
    assert client.post(f"{BASE}/people/{e}/checklists", json={"templateId": inactive["id"]}, headers=hr).status_code == 422
    assert client.post(f"{BASE}/people/{e}/checklists", json={"templateId": 9999}, headers=hr).status_code == 404
    started = client.post(f"{BASE}/people/{e}/checklists", json={"templateId": new["id"], "startedAt": iso(-5)}, headers=hr).json()
    assert started["status"] == "active" and started["startedAt"] == iso(-5)
    first, second = started["items"]
    assert first["dueDate"] == iso(-4) and second["dueDate"] is None and first["assignee"] is None
    # назначение: уведомление исполнителю (не инициатору)
    assigned = client.patch(f"{BASE}/checklist-items/{first['id']}", json={"assigneeUserId": team["users"]["c"].id, "dueDate": iso(7)}, headers=hr).json()
    assert assigned["assignee"] == {"id": team["users"]["c"].id, "name": "Коллегин Клим"} and assigned["dueDate"] == iso(7)
    got = [n for n in notes(client, team["c"]) if n["kind"] == "checklist_assigned"]
    assert len(got) == 1 and got[0]["link"] == "/modules/hr" and "Ермолов Егор" in got[0]["body"]
    assert client.patch(f"{BASE}/checklist-items/{first['id']}", json={"assigneeUserId": 99999}, headers=hr).status_code == 422
    mine = client.get(f"{BASE}/checklists/mine", headers=team["c"]).json()
    assert mine == [{"checklistId": started["id"], "itemId": first["id"], "personName": "Ермолов Егор", "checklistTitle": started["title"], "title": "Аккаунты", "dueDate": iso(7), "done": False}]
    # исполнитель (не HR): только done/note своего пункта
    assert client.patch(f"{BASE}/checklist-items/{first['id']}", json={"dueDate": iso(9)}, headers=team["c"]).status_code == 403
    assert client.patch(f"{BASE}/checklist-items/{first['id']}", json={"assigneeUserId": None}, headers=team["c"]).status_code == 403
    assert client.patch(f"{BASE}/checklist-items/{second['id']}", json={"done": True}, headers=team["c"]).status_code == 404
    assert client.patch(f"{BASE}/checklist-items/{first['id']}", json={"done": True}, headers=team["e"]).status_code == 404
    ok = client.patch(f"{BASE}/checklist-items/{first['id']}", json={"done": True, "note": "готово"}, headers=team["c"]).json()
    assert ok["done"] is True and ok["doneBy"] == "Коллегин Клим" and ok["doneAt"] and ok["note"] == "готово"
    assert client.get(f"{BASE}/people/{e}/checklists", headers=hr).json()[0]["status"] == "active"
    assert not [n for n in notes(client, hr) if n["kind"] == "checklist_done"]
    # последний пункт — чек-лист завершён, уведомление всем HR
    client.patch(f"{BASE}/checklist-items/{second['id']}", json={"done": True}, headers=hr)
    assert client.get(f"{BASE}/people/{e}/checklists", headers=hr).json()[0]["status"] == "done"
    for who in (hr, resp):
        assert len([n for n in notes(client, who) if n["kind"] == "checklist_done"]) == 1
    assert client.patch(f"{BASE}/checklist-items/{second['id']}", json={"done": False}, headers=hr).status_code == 409
    assert client.post(f"{BASE}/checklists/{started['id']}/close", headers=hr).status_code == 409
    assert client.get(f"{BASE}/checklists/mine", headers=team["c"]).json() == []
    assert all(set(d) <= {"checklist_id", "person_id", "kind", "template_id", "item_id", "fields"} for a in ("hr.checklist_start", "hr.checklist_item") for d in audit(db, a))


def test_checklist_close_and_cancel(client, hr, team):
    e = team["ids"]["e"]
    tid = tpl_id(client, hr, "offboarding")
    a = client.post(f"{BASE}/people/{e}/checklists", json={"templateId": tid}, headers=hr).json()
    b = client.post(f"{BASE}/people/{e}/checklists", json={"templateId": tid}, headers=hr).json()
    assert client.post(f"{BASE}/checklists/{a['id']}/close", headers=hr).json()["status"] == "done"
    assert client.post(f"{BASE}/checklists/{b['id']}/cancel", headers=hr).json()["status"] == "cancelled"
    assert client.post(f"{BASE}/checklists/{b['id']}/cancel", headers=hr).status_code == 409
    assert client.post(f"{BASE}/checklists/9999/close", headers=hr).status_code == 404
    assert client.post(f"{BASE}/checklists/{a['id']}/close", headers=team["e"]).status_code == 403


# ---------- график отпусков ----------


def plan(client, headers, start=10, end=16, **kw):
    return client.post(f"{BASE}/me/leave/plans", json={"start": iso(start), "end": iso(end), **kw}, headers=headers)


def test_plan_create_cancel_and_validation(client, db, team):
    h = team["e"]
    created = plan(client, h, comment="Море").json()
    assert created["status"] == "planned" and created["requestId"] is None and created["comment"] == "Море"
    assert plan(client, h, 12, 20).status_code == 409  # пересечение с planned
    assert plan(client, h, 30, 35).status_code == 200
    assert plan(client, h, -5, 3).status_code == 422  # прошлое
    assert client.post(f"{BASE}/me/leave/plans", json={"start": iso(10), "end": iso(5)}, headers=h).status_code == 422
    assert plan(client, h, 800, 805).status_code == 422  # дальше следующего года
    assert plan(client, h, 50, 52, comment="я" * 301).status_code == 422
    year = TODAY.year
    assert len(client.get(f"{BASE}/me/leave/plans?year={year}", headers=h).json()) >= 1
    assert client.get(f"{BASE}/me/leave/plans?year=1800", headers=h).status_code == 422
    assert client.post(f"{BASE}/me/leave/plans/{created['id']}/cancel", headers=h).json()["status"] == "cancelled"
    assert plan(client, h, 11, 14).status_code == 200  # после отмены пересечения нет
    assert client.post(f"{BASE}/me/leave/plans/{created['id']}/cancel", headers=h).status_code == 409
    assert client.post(f"{BASE}/me/leave/plans/{created['id']}/cancel", headers=team["c"]).status_code == 404  # чужой план
    assert all(set(d) == {"plan_id", "person_id"} for d in audit(db, "hr.plan_add"))


def test_plan_without_card_is_rejected(client, outsider):
    assert plan(client, outsider).status_code == 422


def test_plan_convert_to_request(client, db, team):
    h = team["e"]
    p = plan(client, h).json()
    r = client.post(f"{BASE}/me/leave/plans/{p['id']}/request", json={"message": "На связи"}, headers=h)
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["status"] == "converted" and body["requestId"] == body["request"]["id"]
    assert body["request"]["status"] == "pending" and body["request"]["kind"] == "annual" and body["request"]["start"] == iso(10)
    assert body["request"]["message"] == "На связи"
    # обычный поток: заявка у руководителя
    assert [a["id"] for a in client.get(f"{BASE}/me/approvals", headers=team["m"]).json() if a["status"] == "pending"] == [body["requestId"]]
    assert client.post(f"{BASE}/me/leave/plans/{p['id']}/request", json={}, headers=h).status_code == 409
    assert client.post(f"{BASE}/me/leave/plans/{p['id']}/cancel", headers=h).status_code == 409
    # заявка на прошедшие даты — 422, план остаётся planned
    past = db.execute(
        text(
            "INSERT INTO trendsetter_hub.hr_leave_plans (person_id, start_date, end_date) VALUES (:p, :s, :e) RETURNING id"
        ),
        {"p": team["ids"]["e"], "s": TODAY - timedelta(days=9), "e": TODAY - timedelta(days=3)},
    ).scalar_one()
    db.commit()
    assert client.post(f"{BASE}/me/leave/plans/{past}/request", json={}, headers=h).status_code == 422
    assert client.post(f"{BASE}/me/leave/plans/{past}/cancel", headers=h).json()["status"] == "cancelled"


def test_plan_visibility_manager_and_hr(client, hr, team):
    plan(client, team["e"], comment="Егор")
    plan(client, team["c"], 40, 45)
    cancelled = plan(client, team["c"], 60, 62).json()
    client.post(f"{BASE}/me/leave/plans/{cancelled['id']}/cancel", headers=team["c"])
    year = date.fromisoformat(iso(10)).year
    seen = client.get(f"{BASE}/me/leave/team-plan?year={year}", headers=team["m"]).json()
    assert {s["personId"] for s in seen} == {team["ids"]["e"], team["ids"]["c"]} and all(s["status"] == "planned" for s in seen)
    assert client.get(f"{BASE}/me/leave/team-plan?year={year}", headers=team["e"]).status_code == 403  # не руководитель
    assert client.get(f"{BASE}/me/leave/team-plan", headers=hr).status_code == 403  # HR без подчинённых — не Mgr
    all_plans = client.get(f"{BASE}/leave/plan?year={year}", headers=hr).json()
    assert len(all_plans) == 2 and {p["directionId"] for p in all_plans} == {"ops"} and set(all_plans[0]) == {"id", "personId", "name", "directionId", "start", "end", "comment", "status"}


def test_team_plan_requires_linked_manager_card(client, db, team):
    """Mgr — только пользователь, привязанный к карточке (user_id); совпадение по почте или имени права не даёт."""
    m = team["ids"]["m"]
    db.execute(text("UPDATE trendsetter_hub.team_people SET user_id = NULL, email = 'm@example.com' WHERE id = :i"), {"i": m})
    db.commit()
    assert client.get(f"{BASE}/me/leave/team-plan?year={TODAY.year}", headers=team["m"]).status_code == 403


def test_checklist_overview(client, hr, team, viewer):
    e, c = team["ids"]["e"], team["ids"]["c"]
    tid = tpl_id(client, hr)
    old = client.post(f"{BASE}/people/{e}/checklists", json={"templateId": tid, "startedAt": iso(-10)}, headers=hr).json()
    new = client.post(f"{BASE}/people/{c}/checklists", json={"templateId": tid, "startedAt": iso(0)}, headers=hr).json()
    done = client.post(f"{BASE}/people/{c}/checklists", json={"templateId": tid}, headers=hr).json()
    client.post(f"{BASE}/checklists/{done['id']}/close", headers=hr)
    client.patch(f"{BASE}/checklist-items/{old['items'][0]['id']}", json={"done": True}, headers=hr)
    got = client.get(f"{BASE}/checklists", headers=hr).json()
    assert [g["id"] for g in got] == [old["id"], new["id"]]  # сначала с просрочкой
    assert set(got[0]) == {"id", "personId", "personName", "kind", "title", "status", "startedAt", "progress", "overdue"}
    assert got[0]["personName"] == "Ермолов Егор" and got[0]["progress"] == {"done": 1, "total": len(old["items"])}
    assert got[0]["overdue"] > 0 and got[1]["overdue"] == 0
    assert [g["id"] for g in client.get(f"{BASE}/checklists?status=done", headers=hr).json()] == [done["id"]]
    assert len(client.get(f"{BASE}/checklists?status=all", headers=hr).json()) == 3
    assert client.get(f"{BASE}/checklists?status=bad", headers=hr).status_code == 422
    assert client.get(f"{BASE}/checklists", headers=viewer).status_code == 403
    assert client.get(f"{BASE}/checklists", headers=team["e"]).status_code == 403


def test_template_days_offset_bounds(client, hr):
    body = {"kind": "onboarding", "title": "Т"}
    assert client.post(f"{BASE}/checklist-templates", json={**body, "items": [{"title": "П", "daysOffset": 365}]}, headers=hr).status_code == 200
    for bad in (-1, 366):
        assert client.post(f"{BASE}/checklist-templates", json={**body, "items": [{"title": "П", "daysOffset": bad}]}, headers=hr).status_code == 422


def test_history_on_reassign_of_reports(client, hr, team):
    m, e, c = team["ids"]["m"], team["ids"]["e"], team["ids"]["c"]
    ceo = next(p for p in client.get(f"{TEAM}/hr/people", headers=hr).json() if p["isCeo"])["id"]
    assert client.patch(f"{TEAM}/hr/people/{m}", json={"isArchived": True, "reassignTo": ceo}, headers=hr).status_code == 200
    for pid in (e, c):
        found = [h for h in client.get(f"{BASE}/people/{pid}/history", headers=hr).json() if h["kind"] == "manager"]
        assert len(found) == 1 and found[0]["auto"] and found[0]["oldValue"] == "Мэнэджер Мария" and found[0]["newValue"] == "Босс Главный"


# ---------- экспорт ----------


def sheet_text(data: dict) -> dict[str, list[list]]:
    wb = load_workbook(io.BytesIO(base64.b64decode(data["dataBase64"])))
    return {ws.title: [[c.value for c in row] for row in ws.iter_rows()] for ws in wb.worksheets}


def test_export_people_has_no_money_or_documents(client, db, hr, resp, team):
    e = team["ids"]["e"]
    assert client.patch(f"{BASE}/people/{e}/sensitive", json={"salary": SALARY, "passport": SECRET_PASSPORT, "taxId": "7700123456", "bankDetails": "СЧЁТ-777"}, headers=resp).status_code == 200
    upload(client, resp, e, title=SECRET_DOC_TITLE, fileName="секрет.pdf")
    client.patch(f"{BASE}/people/{e}/file", json={"hiredAt": "2024-03-01", "country": "AM", "legalName": "Ермолов Е. Е.", "notes": "заметка"}, headers=resp)
    r = client.get(f"{BASE}/export/people", headers=hr)
    assert r.status_code == 200
    data = r.json()
    assert data["fileName"].endswith(".xlsx") and set(data) == {"fileName", "dataBase64"}
    sheets = sheet_text(data)
    assert list(sheets) == ["Сотрудники"]
    header, *rows = sheets["Сотрудники"]
    assert header == ["ФИО", "Должность", "Направление", "Руководитель", "Статус", "Страна", "Дата приёма", "Заполнено, %", "Почта", "Telegram", "Локация", "Формат работы"]
    egor = next(row for row in rows if row[0] == "Ермолов Егор")
    assert egor[3] == "Мэнэджер Мария" and egor[4] == "Работает" and egor[5] == "Армения" and egor[6] == "01.03.2024" and 0 <= egor[7] <= 100
    blob = " ".join(str(c) for rows_ in sheets.values() for row in rows_ for c in row)
    for secret in ("987654", SECRET_PASSPORT, "7700123456", "СЧЁТ-777", SECRET_DOC_TITLE, "секрет.pdf", "заметка", "Ермолов Е. Е."):
        assert secret not in blob
    import zipfile

    with zipfile.ZipFile(io.BytesIO(base64.b64decode(data["dataBase64"]))) as z:  # и в сыром xml файла
        raw = " ".join(z.read(n).decode("utf-8") for n in z.namelist() if n.endswith(".xml"))
    for secret in ("987654", "СЧЁТ-777", SECRET_DOC_TITLE):
        assert secret not in raw
    assert audit(db, "hr.export") == [{"kind": "people", "rows": len(rows)}]


def test_export_marks_archived_keeps_telegram_and_does_not_make_formulas(client, db, hr, team):
    c, e = team["ids"]["c"], team["ids"]["e"]
    db.execute(text("UPDATE trendsetter_hub.team_people SET is_archived = true, title = '=SUM(A1)' WHERE id = :i"), {"i": c})
    db.execute(text("UPDATE trendsetter_hub.team_people SET answers = '{\"telegram\": \"@ivan\"}'::jsonb WHERE id = :i"), {"i": e})
    db.commit()
    data = client.get(f"{BASE}/export/people", headers=hr).json()
    rows = sheet_text(data)["Сотрудники"]
    klim = next(r for r in rows if r[0] == "Коллегин Клим")
    assert klim[4] == "Уволен" and klim[1] == "=SUM(A1)"
    assert next(r for r in rows if r[0] == "Ермолов Егор")[9] == "@ivan"
    ws = load_workbook(io.BytesIO(base64.b64decode(data["dataBase64"])))["Сотрудники"]
    assert next(row[1] for row in ws.iter_rows(min_row=2) if row[0].value == "Коллегин Клим").data_type == "s"  # не формула


def test_export_leave_and_balance(client, db, hr, team):
    e = team["ids"]["e"]
    client.patch(f"{BASE}/people/{e}/file", json={"hiredAt": (TODAY - timedelta(days=400)).isoformat(), "country": "AM"}, headers=hr)
    assert client.post(f"{BASE}/people/{e}/leave/requests", json={"kind": "annual", "start": iso(5), "end": iso(9), "status": "approved"}, headers=hr).status_code == 200
    client.post(f"{BASE}/people/{e}/leave/requests", json={"kind": "annual", "start": iso(100), "end": iso(101)}, headers=hr)  # pending: не попадёт
    sheets = sheet_text(client.get(f"{BASE}/export/leave?from={iso(0)}&to={iso(30)}", headers=hr).json())
    assert list(sheets) == ["Отпуска", "Баланс"]
    assert sheets["Отпуска"][0] == ["ФИО", "Вид", "С", "По", "Дней"]
    assert sheets["Отпуска"][1:] == [["Ермолов Егор", "Ежегодный", f"{TODAY + timedelta(days=5):%d.%m.%Y}", f"{TODAY + timedelta(days=9):%d.%m.%Y}", 5.0]]
    balance = {r[0]: r for r in sheets["Баланс"][1:]}
    assert sheets["Баланс"][0] == ["ФИО", "Накоплено", "Использовано", "Запланировано", "Доступно"]
    assert balance["Ермолов Егор"][3] == 5.0 and balance["Коллегин Клим"][1] is None  # без даты приёма баланса нет
    assert client.get(f"{BASE}/export/leave?from={iso(0)}&to={iso(401)}", headers=hr).status_code == 422
    assert client.get(f"{BASE}/export/leave?from={iso(10)}&to={iso(0)}", headers=hr).status_code == 422
    assert client.get(f"{BASE}/export/leave?from=bad&to={iso(0)}", headers=hr).status_code == 422
    assert audit(db, "hr.export")[-1]["kind"] == "leave"


# ---------- напоминания и daily_job ----------


def hr_notes(client, who, kind):
    return [n for n in notes(client, who) if n["kind"] == kind]


def test_daily_job_runs_once_per_day_and_dedupes(client, db, hr, team):
    e, c = team["ids"]["e"], team["ids"]["c"]
    client.patch(f"{BASE}/people/{e}/file", json={"probationUntil": iso(3), "hiredAt": f"{TODAY.year - 3}-{TODAY:%m-%d}", "birthDate": f"1990-{TODAY:%m-%d}"}, headers=hr)
    client.patch(f"{BASE}/people/{c}/file", json={"probationUntil": iso(5)}, headers=hr)  # не 14/3/0 — напоминания нет
    assert client.get("/core/me", headers=hr).status_code == 200
    assert client.get("/core/me", headers=hr).status_code == 200  # двойной вызов
    assert len(hr_notes(client, hr, "probation_ending")) == 1
    assert hr_notes(client, hr, "probation_ending")[0]["body"] == f"Скоро заканчивается испытательный срок: Ермолов Егор ({(TODAY + timedelta(days=3)):%d.%m.%Y})"
    assert len(hr_notes(client, hr, "work_anniversary")) == 1 and "3 лет" in hr_notes(client, hr, "work_anniversary")[0]["body"]
    assert len(hr_notes(client, hr, "birthday")) == 1 and "Сегодня день рождения: Ермолов Егор" in hr_notes(client, hr, "birthday")[0]["body"]
    assert len(hr_notes(client, team["m"], "birthday")) == 1  # руководитель получает день рождения подчинённого
    assert hr_notes(client, team["m"], "probation_ending") == [] and hr_notes(client, team["e"], "birthday") == []
    assert hr_notes(client, hr, "probation_ending")[0]["link"] == "/modules/hr"
    # повторный запуск при удалённой отметке дня: события не дублируются (ON CONFLICT по событию)
    from app.modules.hr import reminders

    db.execute(text("DELETE FROM trendsetter_hub.hr_reminder_log WHERE kind = '_run'"))
    db.commit()
    reminders.run_daily(db)
    assert len(hr_notes(client, hr, "probation_ending")) == 1 and len(hr_notes(client, hr, "birthday")) == 1
    reminders.run_daily(db)  # отметка дня есть — ничего не делает
    assert [d for d in audit(db, "hr.reminders_run")][0] == {"probation_ending": 1, "work_anniversary": 1, "birthday": 1, "equipment_return_due": 0, "checklist_overdue": 0}
    # следующий день — снова можно (3 дня до конца уже стало 2: напоминания нет, но запуск проходит)
    reminders.run_daily(db, TODAY + timedelta(days=1))
    assert len(hr_notes(client, hr, "probation_ending")) == 1


def test_daily_job_equipment_and_checklist(client, db, hr, team):
    e = team["ids"]["e"]
    for due in (3, 0, 7):
        client.post(f"{BASE}/people/{e}/equipment", json={"kind": "laptop", "name": f"Ноут {due}", "issuedAt": iso(-30), "returnDue": iso(due)}, headers=hr)
    returned = client.post(f"{BASE}/people/{e}/equipment", json={"kind": "phone", "name": "Вернули", "issuedAt": iso(-30), "returnDue": iso(0), "comment": "x"}, headers=hr).json()
    client.patch(f"{BASE}/equipment/{returned['id']}", json={"returnedAt": iso(-1)}, headers=hr)
    cl = client.post(f"{BASE}/people/{e}/checklists", json={"templateId": tpl_id(client, hr), "startedAt": iso(-10)}, headers=hr).json()
    assert client.get("/core/me", headers=hr).status_code == 200
    eq = hr_notes(client, hr, "equipment_return_due")
    assert sorted(n["body"] for n in eq) == sorted([f"Срок возврата оборудования: Ермолов Егор — Ноут 3 ({(TODAY + timedelta(days=3)):%d.%m.%Y})", f"Срок возврата оборудования: Ермолов Егор — Ноут 0 ({TODAY:%d.%m.%Y})"])
    overdue = hr_notes(client, hr, "checklist_overdue")
    assert len(overdue) == 1 and "Просрочен пункт чек-листа:" in overdue[0]["body"] and "Ермолов Егор" in overdue[0]["body"]
    assert cl["items"][0]["dueDate"] < iso(0)


def test_daily_job_failure_does_not_break_me(client, db, hr, monkeypatch):
    from app.core import routes_core
    from app.core.registry import ModuleManifest

    def boom(session):
        session.execute(text("SELECT 1/0"))

    monkeypatch.setattr(routes_core, "MODULES", [ModuleManifest("hr", "HR", "roadmap", daily_job=boom)])
    r = client.get("/core/me", headers=hr)
    assert r.status_code == 200 and r.json()["unread_notifications"] == 0
    assert client.get("/core/me", headers=hr).status_code == 200  # сессия после сбоя рабочая


def test_daily_job_skipped_when_lock_busy(client, db, hr, team, engine):
    from app.core.routes_core import DAILY_JOBS_LOCK

    client.patch(f"{BASE}/people/{team['ids']['e']}/file", json={"probationUntil": iso(0)}, headers=hr)
    with engine.connect() as other:
        other.execute(text("SELECT pg_advisory_lock(:k)"), {"k": DAILY_JOBS_LOCK})
        try:
            assert client.get("/core/me", headers=hr).status_code == 200
            assert hr_notes(client, hr, "probation_ending") == []
        finally:
            other.execute(text("SELECT pg_advisory_unlock(:k)"), {"k": DAILY_JOBS_LOCK})
            other.commit()
    client.get("/core/me", headers=hr)
    assert len(hr_notes(client, hr, "probation_ending")) == 1


def test_journal_has_no_personal_data(client, db, hr, resp, team):
    e = team["ids"]["e"]
    client.post(f"{BASE}/people/{e}/history", json=history_body(title="Секретный перевод", newValue="СЕКРЕТ"), headers=hr)
    client.post(f"{BASE}/people/{e}/equipment", json={"kind": "laptop", "name": "СЕКРЕТНЫЙ ноут", "serial": "SN-СЕКРЕТ", "issuedAt": iso(0)}, headers=hr)
    tid = tpl_id(client, hr)
    cl = client.post(f"{BASE}/people/{e}/checklists", json={"templateId": tid}, headers=hr).json()
    client.patch(f"{BASE}/checklist-items/{cl['items'][0]['id']}", json={"note": "СЕКРЕТ-заметка", "done": True}, headers=hr)
    client.patch(f"{BASE}/checklist-templates/{tid}", json={"title": "СЕКРЕТ-шаблон"}, headers=hr)
    plan(client, team["e"], comment="СЕКРЕТ-коммент")
    client.get(f"{BASE}/export/people", headers=hr)
    rows = db.execute(text("SELECT action, details FROM trendsetter_hub.audit_log WHERE action LIKE 'hr.%'")).all()
    assert {a for a, _ in rows} >= {"hr.history_add", "hr.equipment_add", "hr.checklist_start", "hr.checklist_item", "hr.template_update", "hr.plan_add", "hr.export"}
    assert "СЕКРЕТ" not in str(rows) and "Секретный" not in str(rows)
