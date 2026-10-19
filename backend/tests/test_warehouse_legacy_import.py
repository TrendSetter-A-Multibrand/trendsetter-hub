"""Разовый перенос Склада из xlsx (app.modules.warehouse.legacy_import). Файлы — только синтетические, в tmp_path."""
import io
import json
import re
import sys
from datetime import datetime, timedelta, timezone
from types import SimpleNamespace
from zoneinfo import ZoneInfo

import openpyxl
import pytest
from sqlalchemy import text

from app.modules.warehouse import legacy_import as li
from app.modules.warehouse import service
from tests.conftest import make_user

T0 = datetime(2026, 10, 14, 9, 0, 0, tzinfo=timezone.utc)  # 12:00 по Москве
KI = [f"010460000000000121" + f"5TEST{n}" for n in range(1, 6)]  # 01 + GTIN + 21 + серия
FULL = [k + "91EE06" for k in KI]
RECEIPT_A, RECEIPT_B = "Поставка А", "Поставка Б"
PLAN_JSON = json.dumps({"name": RECEIPT_A, "map": {"4600000000001": "1", "4600000000002": "3"}, "kiMap": {}}, ensure_ascii=False)
TABLES = ["receipts", "receipt_plans", "scans", "boxes", "storage_boxes", "storage_archive", "extra_sections", "marks", "box_counters", "pallet_pointers"]

HEADERS = {
    "Сводка": ["Приемка", "Отсканировано", "Нет в файле", "Излишек", "Коробов закрыто", "Начало", "Последний скан", "Коробов без палета", "Палетов хранения", "Кодов ЧЗ", "Маркированных без ЧЗ"],
    "Раскладка": ["Штрихкод", "Коробка", "Количество"],
    "Лист 1": [],
    "Сканы": ["ID", "Время", "Сотрудник", "Коробка", "ШК короба", "Баркод", "Статус", "Кто", "Приемка", "КИ Честный знак", "Код ЧЗ полностью", "Палет"],
    "Раскладки": ["Приемка", "Версия", "Часть", "Данные", "Файл", "Загрузил", "Когда"],
    "Короба": ["ШК короба", "Сотрудник", "Коробка", "Товаров", "Открыт", "Закрыт", "Кто", "Приемка", "Палет"],
    "Счётчики": ["Счётчик", "Последний выданный номер"],
    "Приемки": ["Приемка", "Создана", "Кто создал", "Статус"],
    "Хранение": ["Палет", "ШК короба", "Секция", "Палет №", "Штук в коробе", "Поставлен", "Сотрудник", "Кто", "Приемка", "Категория"],
    "Хранение — архив": ["Палет", "ШК короба", "Секция", "Палет №", "Штук в коробе", "Поставлен", "Сотрудник", "Кто", "Приемка", "Категория", "Освобождён", "Кто освободил"],
    "Секции": ["Категория", "Доп. секция", "Назначена", "Кто"],
    "Коды ЧЗ": ["КИ", "GTIN", "Время", "Сотрудник", "Кто", "Приемка", "Коробка", "Код полностью", "Баркод"],
}
METRIC_COLS = {
    "scanned": "Отсканировано", "not_in_file": "Нет в файле", "surplus": "Излишек", "boxes_closed": "Коробов закрыто",
    "boxes_without_pallet": "Коробов без палета", "pallets": "Палетов хранения", "cz_codes": "Кодов ЧЗ", "marked_without_cz": "Маркированных без ЧЗ",
}


def ms(t):
    return int(t.timestamp() * 1000)


def base_data(tz):
    """Содержимое листов (списки словарей). Типы как в реальном файле: float в числах, строки в «Секция»/«Категория»."""
    def loc(t):
        return t.astimezone(ZoneInfo(tz)).replace(tzinfo=None)

    t = [T0 + timedelta(minutes=m) for m in range(6)]
    scan = lambda i, n, box, cat, code, status, who, rec, k: {  # noqa: E731
        "ID": i, "Время": loc(t[n]), "Сотрудник": "ПК-1", "Коробка": cat, "ШК короба": code, "Баркод": f"46000000000{n}", "Статус": status, "Кто": who,
        "Приемка": rec, "КИ Честный знак": KI[k] if k is not None else None, "Код ЧЗ полностью": FULL[k] if k is not None else None, "Палет": "S01_P1",
    }
    d = {name: [] for name in HEADERS}
    d["Приемки"] = [
        {"Приемка": RECEIPT_A, "Создана": loc(t[0]), "Кто создал": "wh1@example.com", "Статус": "открыта"},
        {"Приемка": RECEIPT_B, "Создана": loc(t[0]), "Кто создал": "wh1@example.com", "Статус": "завершена"},
    ]
    d["Сканы"] = [
        scan(f"W-{ms(t[0])}-aaa", 0, 0, "1", "S01_K01_P1_0001", "в файле", "wh1@example.com", RECEIPT_A, 0),
        scan(f"W-{ms(t[1])}-bbb", 1, 0, "1", "S01_K01_P1_0001", "в файле, без ЧЗ", "WH2@Example.com", RECEIPT_A, None),
        scan(f"W-{ms(t[2])}-ccc", 2, 0, "1", "01_0010", "нет в файле", "wh1@example.com", RECEIPT_A, 1),
        scan("IMP-1-0", 3, 0, "3", "S03_K03_P1_0002", "излишек", "wh1@example.com", RECEIPT_A, 2),
        scan(f"W-{ms(t[4])}-ddd", 4, 0, "1", "S01_K01_P1_0006", "нет в файле (чужой код ЧЗ)", "wh1@example.com", RECEIPT_B, 3),
    ]
    box = lambda code, cat, rec: {  # noqa: E731
        "ШК короба": code, "Сотрудник": "ПК-1", "Коробка": cat, "Товаров": 2.0, "Открыт": loc(t[0]), "Закрыт": loc(t[1]), "Кто": "wh1@example.com", "Приемка": rec, "Палет": "S01_P1",
    }
    d["Короба"] = [box("S01_K01_P1_0001", "1", RECEIPT_A), box("S03_K03_P1_0002", "3", RECEIPT_A), box("S01_K01_P1_0006", "1", RECEIPT_B)]
    stor = lambda pal, code, sec, n, cat, rec: {  # noqa: E731
        "Палет": pal, "ШК короба": code, "Секция": sec, "Палет №": float(n), "Штук в коробе": 2.0, "Поставлен": loc(t[1]), "Сотрудник": "ПК-1", "Кто": "wh1@example.com",
        "Приемка": rec, "Категория": cat,
    }
    d["Хранение"] = [stor("S01_P1", "S01_K01_P1_0001", "01", 1, "01", RECEIPT_A), stor("S03_P1", "S03_K03_P1_0002", "03", 1, "03", RECEIPT_A)]
    d["Хранение — архив"] = [{**stor("S01_P1", "S01_K01_P1_0004", "01", 1, "01", RECEIPT_A), "Освобождён": loc(t[5]), "Кто освободил": "wh1@example.com"}]
    d["Секции"] = [{"Категория": 3.0, "Доп. секция": 15.0, "Назначена": loc(t[1]), "Кто": "wh1@example.com"}]
    d["Коды ЧЗ"] = [
        {"КИ": KI[k], "GTIN": KI[k][2:16], "Время": loc(t[n]), "Сотрудник": "ПК-1", "Кто": "wh1@example.com", "Приемка": rec, "Коробка": "1", "Код полностью": FULL[k], "Баркод": f"46000000000{n}"}
        for k, n, rec in [(0, 0, RECEIPT_A), (1, 2, RECEIPT_A), (2, 3, RECEIPT_A), (3, 4, RECEIPT_B)]
    ]
    d["Счётчики"] = [{"Счётчик": c, "Последний выданный номер": v} for c, v in [("BOX_01", 3.0), ("BOX_03", 9.0), ("PAL_01", 1.0), ("PAL_03", 1.0)]]
    version, uploaded = ms(T0 + timedelta(minutes=10)), loc(T0 + timedelta(minutes=10))
    cut = [PLAN_JSON[:20], PLAN_JSON[20:40], PLAN_JSON[40:]]
    d["Раскладки"] = [  # части идут вперемешку
        {"Приемка": rec, "Версия": float(version), "Часть": float(i), "Данные": cut[i], "Файл": "plan.xlsx", "Загрузил": "Иван / wh1@example.com", "Когда": uploaded}
        for rec in (RECEIPT_A, RECEIPT_B) for i in (2, 0, 1)
    ]
    return d


def summary_rows(d):
    """Сводка = те же метрики по сырым строкам листов (с дублями)."""
    out = []
    for r in d["Приемки"]:
        name = r["Приемка"]
        sc = [s for s in d["Сканы"] if s["Приемка"] == name]
        bx = [b for b in d["Короба"] if b["Приемка"] == name]
        stored = {s["ШК короба"] for s in d["Хранение"]}
        vals = {
            "scanned": len(sc), "not_in_file": sum(s["Статус"].startswith("нет в файле") for s in sc), "surplus": sum(s["Статус"].startswith("излишек") for s in sc),
            "boxes_closed": len(bx), "boxes_without_pallet": sum(b["ШК короба"] not in stored for b in bx),
            "pallets": len({s["Палет"] for s in d["Хранение"] if s["Приемка"] == name}), "cz_codes": sum(bool(s["КИ Честный знак"]) for s in sc),
            "marked_without_cz": sum(s["Статус"].endswith("без ЧЗ") for s in sc),
        }
        out.append({"Приемка": name, "Начало": r["Создана"], "Последний скан": r["Создана"], **{METRIC_COLS[k]: float(v) if k == "scanned" else v for k, v in vals.items()}})
    return out


def make_xlsx(tmp_path, tz=li.SHEET_TZ, edit=None, edit_summary=None, name="sklad.xlsx"):
    d = base_data(tz)
    if edit:
        edit(d)
    d["Сводка"] = summary_rows(d)
    if edit_summary:
        edit_summary(d["Сводка"])
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sheet, header in HEADERS.items():
        ws = wb.create_sheet(sheet)
        ws.append(header)
        for row in d[sheet]:
            ws.append([row.get(h) for h in header])
    path = tmp_path / name
    wb.save(path)
    return path


@pytest.fixture()
def users(db):
    make_user(db, "wh1@example.com", "employee")
    make_user(db, "wh2@example.com", "employee", active=False)  # неактивные тоже подходят


def count(db, table):
    return db.execute(text(f"SELECT count(*) FROM trendsetter_hub.warehouse_{table}")).scalar_one()


def counts(db):
    return {t: count(db, t) for t in TABLES}


def assert_empty(db):
    assert counts(db) == dict.fromkeys(TABLES, 0)
    assert db.execute(text("SELECT count(*) FROM trendsetter_hub.audit_log")).scalar_one() == 0


def run(db, path, apply=True, tz=li.SHEET_TZ):
    rep = li.import_workbook(db, path, apply=apply, tz=tz)
    return rep, rep.text()


def audit_count(db, action=li.ACTION):
    return db.execute(text("SELECT count(*) FROM trendsetter_hub.audit_log WHERE action = :a"), {"a": action}).scalar_one()


# ---------- 1. полный перенос ----------


def test_full_import(db, users, tmp_path):
    rep, out = run(db, make_xlsx(tmp_path))
    assert rep.exit_code == 0, out
    assert counts(db) == {
        "receipts": 2, "receipt_plans": 1, "scans": 5, "boxes": 3, "storage_boxes": 2, "storage_archive": 1, "extra_sections": 1, "marks": 4,
        "box_counters": 2, "pallet_pointers": 2,
    }
    # 12:00 по Москве = 09:00 UTC
    at = db.execute(text("SELECT scanned_at FROM trendsetter_hub.warehouse_scans WHERE id LIKE 'W-%-aaa'")).scalar_one()
    assert at == datetime(2026, 10, 14, 9, 0, tzinfo=timezone.utc)
    # почта без учёта регистра, worker как есть, статусы
    rows = {r.id.split("-")[-1]: r for r in db.execute(text("SELECT s.*, u.email FROM trendsetter_hub.warehouse_scans s LEFT JOIN trendsetter_hub.users u ON u.id = s.user_id"))}
    assert rows["bbb"].email == "wh2@example.com" and rows["bbb"].worker == "ПК-1"
    assert (rows["bbb"].status, rows["bbb"].no_cz, rows["bbb"].ki, rows["bbb"].full_code) == ("ok", True, None, None)
    assert (rows["ccc"].status, rows["ddd"].status, rows["0"].status) == ("unknown", "foreign", "extra")
    assert rows["aaa"].full_code == FULL[0] and rows["aaa"].box_category == "1"
    # секция 3.0 / 15.0 -> «03» / «15»
    assert db.execute(text("SELECT category, section FROM trendsetter_hub.warehouse_extra_sections")).one() == ("03", "15")
    # Раскладка склеена из перемешанных частей; раскладка завершённой приёмки пропущена с предупреждением
    plan = db.execute(text("SELECT r.name, p.version, p.data, p.worker, u.email FROM trendsetter_hub.warehouse_receipt_plans p JOIN trendsetter_hub.warehouse_receipts r ON r.id = p.receipt_id LEFT JOIN trendsetter_hub.users u ON u.id = p.uploaded_by")).one()
    assert plan.name == RECEIPT_A and plan.data == json.loads(PLAN_JSON) and plan.worker == "Иван" and plan.email == "wh1@example.com"
    assert isinstance(plan.version, int) and plan.version == ms(T0 + timedelta(minutes=10))
    assert "завершённой приёмки" in out
    # счётчик = max(лист, номера кодов, включая временный 01_0010)
    assert dict(db.execute(text("SELECT category, last_number FROM trendsetter_hub.warehouse_box_counters")).all()) == {"01": 10, "03": 9}
    # хаб: тип и значение палета/товаров
    assert db.execute(text("SELECT pallet_no, item_count, category FROM trendsetter_hub.warehouse_storage_boxes WHERE section = '01'")).one() == (1, 2, "01")
    # ровно одна запись журнала, без КИ и почт
    assert audit_count(db) == 1
    details = db.execute(text("SELECT details::text FROM trendsetter_hub.audit_log")).scalar_one()
    assert "@" not in details and not re.search(r"01\d{14}21", details) and "file_sha256" in details
    # представление совпадает со Сводкой
    view = {r.receipt: r for r in db.execute(text("SELECT * FROM trendsetter_hub.warehouse_receipt_summary"))}
    assert (view[RECEIPT_A].scanned, view[RECEIPT_A].not_in_file, view[RECEIPT_A].surplus, view[RECEIPT_A].cz_codes, view[RECEIPT_A].marked_without_cz) == (4, 1, 1, 3, 1)
    assert (view[RECEIPT_B].boxes_closed, view[RECEIPT_B].boxes_without_pallet, view[RECEIPT_A].pallets) == (1, 1, 2)
    assert "временных коробах" in out and "01_0010×1" in out


# ---------- 2-3. идемпотентность и сухой прогон ----------


def test_idempotent(db, users, tmp_path):
    path = make_xlsx(tmp_path)
    assert run(db, path)[0].exit_code == 0
    before = counts(db)
    rep, out = run(db, path)
    assert rep.exit_code == 0 and "изменений нет" in out
    assert counts(db) == before and audit_count(db) == 1


def test_dry_run_writes_nothing(db, users, tmp_path):
    rep, out = run(db, make_xlsx(tmp_path), apply=False)
    assert rep.exit_code == 0, out
    assert rep.lines[0].startswith("Перенос «Склад-приемка»: СУХОЙ ПРОГОН (ничего не записано)")
    assert re.search(r"scans\s+5\s+0\s+5\s+0", out) and re.search(r"marks\s+4\s+0\s+4\s+0", out)
    assert_empty(db)


# ---------- 4. неизвестная почта ----------


def test_unknown_email(db, tmp_path):
    make_user(db, "wh1@example.com", "employee")
    rep, out = run(db, make_xlsx(tmp_path))
    assert rep.exit_code == 1 and "wh2@example.com" in out
    assert_empty(db)


# ---------- 5. битые строки ----------


def _set(sheet, i, col, value):
    return lambda d: d[sheet][i].__setitem__(col, value)


def _drop_part(d):
    d["Раскладки"] = [r for r in d["Раскладки"] if not (r["Приемка"] == RECEIPT_A and r["Часть"] == 1.0)]


def _clash(d):
    d["Сканы"].append({**d["Сканы"][0], "Сотрудник": "ДРУГОЙ"})


@pytest.mark.parametrize(
    ("edit", "sheet", "reason"),
    [
        (_set("Сканы", 0, "КИ Честный знак", "123"), "Сканы", "неверный КИ"),
        (_set("Сканы", 0, "Статус", "странный"), "Сканы", "неизвестный статус"),
        (_set("Сканы", 1, "Время", "вчера"), "Сканы", "нечитаемая дата"),
        (_set("Сканы", 1, "Приемка", None), "Сканы", "пусто"),
        (_set("Коды ЧЗ", 0, "Приемка", None), "Коды ЧЗ", "пусто"),
        (_set("Хранение", 0, "Палет", "S01_P2"), "Хранение", "не согласован"),
        (_drop_part, "Раскладки", "без пропусков"),
        (_clash, "Сканы", "другим содержимым"),
    ],
)
def test_broken_rows(db, users, tmp_path, edit, sheet, reason):
    rep, out = run(db, make_xlsx(tmp_path, edit=edit))
    assert rep.exit_code == 1, out
    assert re.search(rf"{sheet} \| строка \d+ \| .* \| .*{reason}", out), out
    assert_empty(db)


def test_missing_header(db, users, tmp_path):
    path = make_xlsx(tmp_path)
    wb = openpyxl.load_workbook(path)
    wb["Сканы"]["F1"] = "Что-то другое"  # был «Баркод»
    wb.save(path)
    rep, out = run(db, path)
    assert rep.exit_code == 1 and "Баркод" in out and "нет обязательного заголовка" in out
    assert_empty(db)


# ---------- 6. конфликты с БД ----------


def _only_file_tables_empty(db, keep):
    c = counts(db)
    assert {t: n for t, n in c.items() if t not in keep} == {t: 0 for t in TABLES if t not in keep}


def test_conflict_foreign_receipt(db, users, tmp_path):
    db.execute(text("INSERT INTO trendsetter_hub.warehouse_receipts (name) VALUES ('Чужая')"))
    db.commit()
    rep, out = run(db, make_xlsx(tmp_path))
    assert rep.exit_code == 2 and "receipts" in out and "Чужая" in out
    _only_file_tables_empty(db, {"receipts"})
    assert count(db, "receipts") == 1


def test_conflict_foreign_mark(db, users, tmp_path):
    rid = db.execute(text("INSERT INTO trendsetter_hub.warehouse_receipts (name) VALUES (:n) RETURNING id"), {"n": RECEIPT_A}).scalar_one()
    foreign = "010460000000000121" + "5FOREIGN"
    db.execute(text("INSERT INTO trendsetter_hub.warehouse_marks (ki, gtin, receipt_id) VALUES (:k, '04600000000001', :r)"), {"k": foreign, "r": rid})
    db.commit()
    rep, out = run(db, make_xlsx(tmp_path))
    assert rep.exit_code == 2 and "marks" in out and foreign not in out and foreign[:16] + "…" in out
    _only_file_tables_empty(db, {"receipts", "marks"})


def test_conflict_counter_ahead_of_file(db, users, tmp_path):
    for _ in range(11):  # в файле BOX_01 расчётно 10
        db.execute(text("SELECT trendsetter_hub.warehouse_next_box_number('01')"))
    db.commit()
    rep, out = run(db, make_xlsx(tmp_path))
    assert rep.exit_code == 2 and "box_counters" in out
    _only_file_tables_empty(db, {"box_counters"})


# ---------- 7. дубли строк «Короба» ----------


def test_duplicate_boxes_collapsed(db, users, tmp_path):
    rep, out = run(db, make_xlsx(tmp_path, edit=lambda d: d["Короба"].append(dict(d["Короба"][0]))))
    assert rep.exit_code == 0, out
    assert count(db, "boxes") == 3 and "дублей строк «Короба»" in out
    assert re.search(r"boxes\s+4\s+1\s+3\s+0", out)
    assert f"{RECEIPT_A} | коробов закрыто | 3 | 3 | 1 | 2 | OK" in out


# ---------- 8. сводка не сходится ----------


def test_summary_mismatch(db, users, tmp_path):
    def bad(rows):
        rows[0]["Отсканировано"] += 1
    rep, out = run(db, make_xlsx(tmp_path, edit_summary=bad))
    assert rep.exit_code == 3 and "РАСХОЖДЕНИЕ со Сводкой" in out
    assert_empty(db)


# ---------- 9. часовой пояс ----------


def test_timezone_selfcheck(db, users, tmp_path):
    path = make_xlsx(tmp_path, tz="Asia/Yerevan")  # время в файле на 1 ч больше московского
    rep, out = run(db, path, tz="Europe/Moscow")
    assert rep.exit_code == 1 and "tz | строка 0 | --tz" in out and "+1 ч" in out
    assert_empty(db)
    rep, out = run(db, path, tz="Asia/Yerevan")
    assert rep.exit_code == 0, out
    assert db.execute(text("SELECT scanned_at FROM trendsetter_hub.warehouse_scans WHERE id LIKE 'W-%-aaa'")).scalar_one() == T0


# ---------- 10. откат ----------


def test_rollback(db, users, tmp_path):
    assert run(db, make_xlsx(tmp_path))[0].exit_code == 0
    full = counts(db)
    rep = li.rollback_import(db, apply=False)
    assert rep.exit_code == 0 and counts(db) == full
    rep = li.rollback_import(db, apply=True)
    assert rep.exit_code == 0, rep.text()
    assert counts(db) == dict.fromkeys(TABLES, 0)
    assert audit_count(db, li.ACTION_ROLLBACK) == 1 and audit_count(db) == 1
    assert li.rollback_import(db, apply=True).exit_code == 2


def test_rollback_refused_after_live_work(db, users, tmp_path):
    assert run(db, make_xlsx(tmp_path))[0].exit_code == 0
    user = db.execute(text("SELECT id FROM trendsetter_hub.users WHERE email = 'wh1@example.com'")).one()
    op = {"op": "scan", "id": "W-1-live", "t": ms(T0), "receipt": RECEIPT_A, "box": "1", "boxCode": "01_0011", "bc": "1", "status": "ok", "pc": "ПК"}
    service.sync(db, SimpleNamespace(id=user.id), [op], "ПК")
    full = counts(db)
    rep = li.rollback_import(db, apply=True)
    assert rep.exit_code == 2 and "живая работа" in rep.text()
    assert counts(db) == full and audit_count(db, li.ACTION_ROLLBACK) == 0


# ---------- 11. маскирование в выводе, main ----------


def test_main_masks_codes(db, users, tmp_path, monkeypatch, capsys, engine):
    monkeypatch.setattr(li, "get_engine", lambda: engine)
    path = make_xlsx(tmp_path)
    assert li.main([str(path)]) == 0
    captured = capsys.readouterr()
    out = captured.out + captured.err
    assert "СУХОЙ ПРОГОН" in out
    assert not any(k in out for k in KI + FULL) and not re.search(r"01\d{14}21\w", out)
    assert "@" not in out
    assert_empty(db)


# ---------- 12. stdin ----------


def test_stdin(db, users, tmp_path, monkeypatch):
    path = make_xlsx(tmp_path)
    expected = run(db, path, apply=False)[1]
    monkeypatch.setattr(sys, "stdin", SimpleNamespace(buffer=io.BytesIO(path.read_bytes())))
    rep, out = run(db, "-", apply=False)
    assert rep.exit_code == 0 and out == expected
    assert_empty(db)
